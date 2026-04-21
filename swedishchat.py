"""
Swedish AI Chat — Flask Backend
════════════════════════════════════════════════════════════════════════
KB-Lab NLP Pipeline (4 automatic stages on every message):

  STAGE 1 | LANGUAGE DETECTION
  STAGE 2 | SPELL-CHECK & BROKEN SENTENCE REPAIR
  STAGE 3 | INTENT & QUESTION UNDERSTANDING
  STAGE 4 | QUERY-FOCUSED EXTRACTIVE SUMMARISATION

  FIXES APPLIED (v2):
  [FIX-1] Stage 4: never summarise replies that contain code blocks,
          numbered lists, or bullet points — return full reply as-is.
  [FIX-2] num_predict raised: 512/1024/768 (was 256/768/512).
  [FIX-3] _split_sentences: only split on sentence-ending punctuation
          followed by an uppercase letter — not on bare newlines.
  [FIX-4] Direct-path max_sents raised to 5 (was 3).
  [FIX-5] Stage 2: skip known tech/code terms to avoid corrupting them.
  [FIX-6] Markdown preserved end-to-end; TTS cleaning only happens
          inside synth_tts(), never mutates the chat reply.

  FILE UPLOAD (v3):
  [UP-1]  /api/chat now accepts multipart/form-data OR application/json.
  [UP-2]  Uploaded files saved to /uploads, text extracted, injected
          into Ollama system prompt under === BIFOGADE FILER ===.
  [UP-3]  Images encoded as base64 data-URIs for vision models.
  [UP-4]  Temp files cleaned up after every response path.

  IMAGE GENERATION (v4):
  [IMG-1] Swedish keyword detection triggers image generation in /api/chat.
  [IMG-2] Prompt auto-translated to English via Ollama before generation.
  [IMG-3] Uses HuggingFace diffusers StableDiffusionPipeline (local).
  [IMG-4] Dedicated /api/generate-image endpoint added.
  [IMG-5] Generated images saved to /generated_images/ and returned as base64.
  [IMG-6] Model: runwayml/stable-diffusion-v1-5 (override via env-var).

  STREAMING & UI FEATURES (v5):
  [STR-1] /api/chat/stream  — SSE streaming endpoint (word-by-word tokens).
  [STR-2] /api/chat/stop    — Cancel an in-progress generation.
  [STR-3] /api/threads/<id>/messages/<idx>/edit   — Edit a user message & resend.
  [STR-4] /api/threads/<id>/regenerate            — Regenerate last assistant reply.
  [STR-5] /api/threads/search                     — Full-text search across threads.

  IMAGE GENERATION FIXES (v6):
  [IMG-FIX-1] is_image_generation_request: Swedish creation verbs only — no
              false positives on generic words like "bild", "foto", "image".
  [IMG-FIX-2] Image generation check runs BEFORE Tavily search so it can
              never be intercepted by the web-search path.
  [IMG-FIX-3] _translate_prompt_to_english: actually enabled and calls Ollama.
  [IMG-FIX-4] English input for image generation is rejected with Swedish-only
              message in both /api/chat and /api/chat/stream.
════════════════════════════════════════════════════════════════════════
"""
from dotenv import load_dotenv
load_dotenv()
import io
import os
import re
import uuid
import json
import wave
import base64
import logging
import mimetypes
import tempfile
import threading
import requests
from datetime import datetime
from collections import defaultdict
from flask import Flask, request, jsonify, session, send_file, Response, stream_with_context
from werkzeug.utils import secure_filename

# ── Core ML imports ───────────────────────────────────────────────────────────
try:
    import torch
    import torch.nn.functional as F
    from transformers import (
        AutoTokenizer,
        AutoModel,
        pipeline as hf_pipeline,
    )
    TRANSFORMERS_OK = True
except ImportError:
    TRANSFORMERS_OK = False

# ── HuggingFace diffusers — Stable Diffusion image generation ─────────────────
try:
    from diffusers import StableDiffusionPipeline
    DIFFUSERS_OK = True
except ImportError:
    DIFFUSERS_OK = False
    StableDiffusionPipeline = None

# ── KBLab Piper TTS ───────────────────────────────────────────────────────────
try:
    from piper.voice import PiperVoice
    PIPER_OK = True
except ImportError:
    PIPER_OK = False

# ── KBLab wav2vec2 STT ────────────────────────────────────────────────────────
try:
    import torchaudio
    from transformers import Wav2Vec2ForCTC, Wav2Vec2Processor
    WAV2VEC_OK = True
except ImportError:
    WAV2VEC_OK = False

# ── Optional: PDF / DOCX / XLSX extraction ────────────────────────────────────
try:
    import pdfplumber
    PDFPLUMBER_OK = True
except ImportError:
    PDFPLUMBER_OK = False

try:
    from docx import Document as DocxDocument
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import openpyxl
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

try:
    from flask_cors import CORS
    HAS_CORS = True
except ImportError:
    HAS_CORS = False

# ── Flask ─────────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET", "swedish-chat-2024")
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

ALLOWED = [
    "http://localhost:4200", "http://127.0.0.1:4200",
    "http://localhost:3000", "http://127.0.0.1:3000",
    "https://pykara.ai"
]

if HAS_CORS:
    CORS(app, supports_credentials=True, origins=ALLOWED)
else:
    @app.after_request
    def _cors(resp):
        origin = request.headers.get("Origin", "")
        if any(origin.startswith(o) for o in ALLOWED):
            resp.headers["Access-Control-Allow-Origin"]      = origin
            resp.headers["Access-Control-Allow-Credentials"] = "true"
            resp.headers["Access-Control-Allow-Headers"]     = "Content-Type,Authorization,X-User-ID"
            resp.headers["Access-Control-Allow-Methods"]     = "GET,POST,DELETE,OPTIONS"
        return resp

    @app.route("/api/<path:p>", methods=["OPTIONS"])
    def _preflight(p):
        return "", 204

# ── Config ────────────────────────────────────────────────────────────────────
OLLAMA_BASE  = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL",    "llama3.2:3b")
TAVILY_KEY   = os.environ.get("TAVILY_API_KEY",  "")

MAX_THREAD_MESSAGES = 40
SUMMARY_TRIGGER     = 30

BERT_ID = "KBLab/bert-base-swedish-cased"
NER_ID  = "KBLab/bert-base-swedish-cased-ner"

PIPER_MODEL_PATH = os.environ.get(
    "PIPER_MODEL_PATH",
    "KBLab/piper-tts-nst-swedish",
)
KB_STT_MODEL_ID = "KBLab/wav2vec2-large-voxrex-swedish"

# ── File upload config ────────────────────────────────────────────────────────
UPLOAD_FOLDER      = os.path.join(os.path.dirname(__file__), "uploads")
MAX_FILE_BYTES     = 20 * 1024 * 1024   # 20 MB per file
MAX_FILES          = 10
ALLOWED_IMAGE_EXTS = {"jpg", "jpeg", "png", "gif", "webp"}
ALLOWED_DOC_EXTS   = {"pdf", "txt", "md", "csv", "json",
                       "doc", "docx", "xls", "xlsx", "pptx", "zip"}
ALLOWED_EXTS       = ALLOWED_IMAGE_EXTS | ALLOWED_DOC_EXTS

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ── Image generation config ───────────────────────────────────────────────────
SD_MODEL_ID          = os.environ.get("SD_MODEL_ID", "runwayml/stable-diffusion-v1-5")
SD_STEPS             = int(os.environ.get("SD_STEPS",    "20"))
SD_CFG_SCALE         = float(os.environ.get("SD_CFG",   "7.0"))
SD_WIDTH             = int(os.environ.get("SD_WIDTH",   "512"))
SD_HEIGHT            = int(os.environ.get("SD_HEIGHT",  "512"))
SD_NEGATIVE_DEFAULT  = os.environ.get(
    "SD_NEGATIVE",
    "blurry, ugly, deformed, low quality, watermark, text",
)
IMGGEN_OUTPUT_FOLDER = os.path.join(os.path.dirname(__file__), "generated_images")
os.makedirs(IMGGEN_OUTPUT_FOLDER, exist_ok=True)

# ── Shared model handles ──────────────────────────────────────────────────────
_tok        = None
_bert       = None
_fill_mask  = None
_ner_pipe   = None
_piper      = None
_stt_proc   = None
_stt_model  = None
_sd_pipe    = None          # [IMG-3] Stable Diffusion pipeline (lazy-loaded)

# ════════════════════════════════════════════════════════════════════════
# [STR-2] Active generation tracking — for stop/cancel support
# ════════════════════════════════════════════════════════════════════════
_active_generations: dict = {}
_active_gen_lock = threading.Lock()


def _get_stop_event(sid: str) -> threading.Event:
    with _active_gen_lock:
        if sid not in _active_generations:
            _active_generations[sid] = threading.Event()
        return _active_generations[sid]


def _reset_stop_event(sid: str) -> threading.Event:
    evt = threading.Event()
    with _active_gen_lock:
        _active_generations[sid] = evt
    return evt


# ════════════════════════════════════════════════════════════════════════
# Thread store
# ════════════════════════════════════════════════════════════════════════
ALL_THREADS:   dict = {}
SESSION_INDEX: dict = {}
ACTIVE_THREAD: dict = {}
USER_IMAGES:   dict = {}   # sid -> list[{id, b64, prompt_sv, prompt_en, created_at}]

PERSIST_FILE = os.environ.get("THREAD_STORE_PATH",
                              os.path.join(os.path.dirname(__file__), "thread_store.json"))
_PERSIST_DIRTY = False


def _persist_save() -> None:
    global _PERSIST_DIRTY
    payload = {
        "all_threads":   ALL_THREADS,
        "session_index": SESSION_INDEX,
        "active_thread": ACTIVE_THREAD,
        "user_images":   USER_IMAGES,
    }
    tmp = PERSIST_FILE + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        os.replace(tmp, PERSIST_FILE)
        _PERSIST_DIRTY = False
        logger.debug(f"[Persist] Saved {len(ALL_THREADS)} thread(s) → {PERSIST_FILE}")
    except Exception as e:
        logger.warning(f"[Persist] Save failed: {e}")


def _persist_load() -> None:
    global ALL_THREADS, SESSION_INDEX, ACTIVE_THREAD, USER_IMAGES
    if not os.path.exists(PERSIST_FILE):
        logger.info(f"[Persist] No store file found at {PERSIST_FILE} — starting fresh")
        return
    try:
        with open(PERSIST_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        loaded_threads = data.get("all_threads", {})
        valid = {}
        required = {"id", "title", "created_at", "updated_at", "messages", "message_count"}
        for tid, t in loaded_threads.items():
            if required.issubset(t.keys()):
                t.setdefault("summary",          None)
                t.setdefault("_last_summary_at", 0)
                valid[tid] = t
            else:
                logger.warning(f"[Persist] Skipping malformed thread {tid}")
        ALL_THREADS   = valid
        SESSION_INDEX = data.get("session_index", {})
        ACTIVE_THREAD = data.get("active_thread", {})
        USER_IMAGES   = data.get("user_images",   {})
        logger.info(f"[Persist] Loaded {len(ALL_THREADS)} thread(s) and {sum(len(v) for v in USER_IMAGES.values())} image(s) from {PERSIST_FILE}")
    except Exception as e:
        logger.error(f"[Persist] Load failed ({e}) — starting with empty store")


def _mark_dirty() -> None:
    global _PERSIST_DIRTY
    _PERSIST_DIRTY = True


def _get_store(sid: str) -> dict:
    tids   = SESSION_INDEX.get(sid, [])
    active = ACTIVE_THREAD.get(sid)
    return {
        "threads":       {tid: ALL_THREADS[tid] for tid in tids if tid in ALL_THREADS},
        "active_thread": active,
    }


def _register_thread(sid: str, thread: dict) -> None:
    ALL_THREADS[thread["id"]] = thread
    if sid not in SESSION_INDEX:
        SESSION_INDEX[sid] = []
    if thread["id"] not in SESSION_INDEX[sid]:
        SESSION_INDEX[sid].append(thread["id"])
    ACTIVE_THREAD[sid] = thread["id"]
    _persist_save()


def _resolve_thread(sid: str, thread_id) -> dict:
    if thread_id and thread_id in ALL_THREADS:
        ACTIVE_THREAD[sid] = thread_id
        if sid not in SESSION_INDEX:
            SESSION_INDEX[sid] = []
        if thread_id not in SESSION_INDEX[sid]:
            SESSION_INDEX[sid].append(thread_id)
        return ALL_THREADS[thread_id]

    active_tid = ACTIVE_THREAD.get(sid)
    if active_tid and active_tid in ALL_THREADS:
        return ALL_THREADS[active_tid]

    thread = _make_thread()
    _register_thread(sid, thread)
    logger.info(f"[Thread] Auto-created thread {thread['id']} for sid={sid[:8]}")
    return thread


# ════════════════════════════════════════════════════════════════════════
# Load all KB-Lab models once at startup
# ════════════════════════════════════════════════════════════════════════
def load_kb_models():
    global _tok, _bert, _fill_mask, _ner_pipe, _piper, _stt_proc, _stt_model

    if not TRANSFORMERS_OK:
        logger.warning("transformers/torch not installed — KB-Lab stages use fallbacks")
        return

    logger.info("Loading KB-Lab models…")

    try:
        _tok = AutoTokenizer.from_pretrained(BERT_ID, use_fast=True)
        logger.info("  [OK] Tokenizer  KBLab/bert-base-swedish-cased")
    except Exception as e:
        logger.warning(f"  [FAIL] Tokenizer: {e}")

    try:
        _fill_mask = hf_pipeline("fill-mask", model=BERT_ID, top_k=5, device=-1)
        logger.info("  [OK] Fill-mask  KBLab/bert-base-swedish-cased  (Stage 2)")
    except Exception as e:
        logger.warning(f"  [FAIL] Fill-mask: {e}")

    try:
        _bert = AutoModel.from_pretrained(BERT_ID)
        _bert.eval()
        logger.info("  [OK] BERT base  KBLab/bert-base-swedish-cased  (Stage 4)")
    except Exception as e:
        logger.warning(f"  [FAIL] BERT base: {e}")

    try:
        _ner_pipe = hf_pipeline(
            "ner", model=NER_ID,
            aggregation_strategy="simple", device=-1,
        )
        logger.info("  [OK] NER        KBLab/bert-base-swedish-cased-ner  (Stage 3)")
    except Exception as e:
        logger.warning(f"  [FAIL] NER: {e}")

    _load_piper()
    _load_wav2vec2()

    logger.info("KB-Lab loading complete")


def _load_piper():
    global _piper
    if not PIPER_OK:
        logger.warning("  [SKIP] piper-tts not installed  (pip install piper-tts)")
        return

    local_path = PIPER_MODEL_PATH if PIPER_MODEL_PATH.endswith(".onnx") else None

    if local_path and os.path.isfile(local_path):
        onnx_path = local_path
    else:
        try:
            from huggingface_hub import hf_hub_download
            onnx_path = hf_hub_download(
                repo_id  = "KBLab/piper-tts-nst-swedish",
                filename = "epoch=4041-step=1753548.onnx",
            )
        except Exception:
            logger.warning(
                "  [WARN] KBLab Piper ONNX not found. "
                "Set PIPER_MODEL_PATH=/path/to/model.onnx"
            )
            return

    try:
        _piper = PiperVoice.load(onnx_path, use_cuda=False)
        logger.info(f"  [OK] Piper TTS  KBLab/piper-tts-nst-swedish  ({onnx_path})")
    except Exception as e:
        logger.warning(f"  [FAIL] Piper TTS: {e}")


def _load_wav2vec2():
    global _stt_proc, _stt_model
    if not WAV2VEC_OK:
        logger.warning("  [SKIP] torchaudio/transformers not fully available for STT")
        return
    try:
        _stt_proc  = Wav2Vec2Processor.from_pretrained(KB_STT_MODEL_ID)
        _stt_model = Wav2Vec2ForCTC.from_pretrained(KB_STT_MODEL_ID)
        _stt_model.eval()
        logger.info(f"  [OK] STT  {KB_STT_MODEL_ID}")
    except Exception as e:
        logger.warning(f"  [FAIL] STT wav2vec2: {e}")


# ════════════════════════════════════════════════════════════════════════
# [IMG-3] Stable Diffusion pipeline — lazy loader
# ════════════════════════════════════════════════════════════════════════

def _load_sd_pipeline():
    global _sd_pipe
    if _sd_pipe is not None:
        return _sd_pipe

    if not DIFFUSERS_OK:
        logger.warning(
            "[ImgGen] diffusers not installed. "
            "Run: pip install diffusers transformers accelerate"
        )
        return None

    device = "cuda" if torch.cuda.is_available() else "cpu"
    dtype  = torch.float16 if device == "cuda" else torch.float32

    logger.info(f"[ImgGen] Loading StableDiffusionPipeline: {SD_MODEL_ID} on {device} …")
    try:
        _sd_pipe = StableDiffusionPipeline.from_pretrained(
            SD_MODEL_ID,
            torch_dtype=dtype,
        ).to(device)
        _sd_pipe.safety_checker = None
        logger.info(f"[ImgGen] StableDiffusionPipeline ready on {device}")
    except Exception as e:
        logger.error(f"[ImgGen] Failed to load pipeline: {e}")
        _sd_pipe = None

    return _sd_pipe


# ════════════════════════════════════════════════════════════════════════
# [IMG-FIX-1] Swedish image-generation keyword detection
# Only triggers on explicit Swedish creation verbs.
# Does NOT trigger on bare words like "bild", "foto", "image".
# ════════════════════════════════════════════════════════════════════════

def is_image_generation_request(text: str) -> bool:
    """
    Returns True only when the user is asking to CREATE/GENERATE an image
    using Swedish creation verbs. Generic mentions of 'bild' or 'foto' alone
    do NOT trigger image generation.
    """
    t = text.lower().strip()

    # Must contain an explicit Swedish creation verb phrase
    creation_phrases = [
        "rita en ", "rita ett ", "rita bild", "rita mig",
        "skapa en bild", "skapa ett ", "skapa bild", "skapa illustration",
        "generera en bild", "generera ett ", "generera bild", "generera illustration",
        "måla en ", "måla ett ", "måla bild", "måla mig",
        "illustrera en ", "illustrera ett ",
        "gör en bild", "gör ett ", "gör mig en bild",
        "ge mig en bild", "ge mig ett ",
        "visa mig en bild av", "visa mig ett ",
        "skapa teckning", "rita teckning",
        "generera foto", "skapa foto",
        "bildgenerering", "text till bild", "text-till-bild",
    ]
    if any(phrase in t for phrase in creation_phrases):
        return True

    # "bild av/på/med" — but ONLY combined with creation verbs, not standalone
    if re.search(r"\b(skapa|generera|rita|måla|gör|ge)\b.*\bbild\s+(av|på|med)\b", t):
        return True

    # "kan du rita/skapa/generera/måla ..."
    if re.search(
        r"\bkan\s+du\s+(rita|skapa|generera|måla|illustrera)\b", t
    ):
        return True

    # "snälla rita/skapa ..."
    if re.search(
        r"\bsnälla\s+(rita|skapa|generera|måla|illustrera)\b", t
    ):
        return True

    return False


def extract_image_prompt(text: str) -> str:
    t = text.strip()
    t = re.sub(
        r"^(kan\s+du\s+)?(snälla\s+)?"
        r"(rita|skapa|generera|måla|illustrera|gör(\s+mig)?)\s+"
        r"(en\s+|ett\s+)?",
        "",
        t,
        flags=re.IGNORECASE,
    )
    t = re.sub(r"^(en\s+bild\s+(av|på|med)\s+)", "", t, flags=re.IGNORECASE)
    t = re.sub(r"^(bild\s+(av|på|med)\s+)", "", t, flags=re.IGNORECASE)
    return t.strip() or text.strip()


# ════════════════════════════════════════════════════════════════════════
# [IMG-FIX-3] Translate Swedish prompt → English via Ollama (ENABLED)
# ════════════════════════════════════════════════════════════════════════

def _translate_prompt_to_english(swedish_prompt: str) -> str:
    """Translate a Swedish image prompt to English using Ollama."""
    try:
        payload = {
            "model": OLLAMA_MODEL,
            "messages": [{
                "role": "user",
                "content": (
                    "Translate this Swedish image description to English for use "
                    "as a Stable Diffusion prompt. Return ONLY the English "
                    "translation, nothing else, no explanation:\n\n"
                    + swedish_prompt
                ),
            }],
            "stream": False,
            "options": {"temperature": 0.1, "num_predict": 200},
        }
        r = requests.post(
            f"{OLLAMA_BASE}/api/chat", json=payload, timeout=(5, 30)
        )
        r.raise_for_status()
        translated = r.json().get("message", {}).get("content", "").strip()
        # Strip any wrapping quotes Ollama might add
        translated = translated.strip('"\'')
        if translated:
            logger.info(f"[ImgGen] Translated: '{swedish_prompt}' → '{translated}'")
            return translated
        return swedish_prompt
    except Exception as e:
        logger.warning(f"[ImgGen] Translation failed ({e}) — using original prompt")
        return swedish_prompt


# ════════════════════════════════════════════════════════════════════════
# [IMG-3/5] Core image generation — diffusers StableDiffusionPipeline
# ════════════════════════════════════════════════════════════════════════

def _save_generated_image(image_bytes: bytes) -> str:
    fname = f"img_{uuid.uuid4().hex}.png"
    path  = os.path.join(IMGGEN_OUTPUT_FOLDER, fname)
    try:
        with open(path, "wb") as f:
            f.write(image_bytes)
        logger.info(f"[ImgGen] Saved → {path}")
    except Exception as e:
        logger.warning(f"[ImgGen] Could not save image file: {e}")
    return path


def generate_image_diffusers(
    prompt:          str,
    negative_prompt: str = "",
) -> dict:
    pipe = _load_sd_pipeline()
    if pipe is None:
        return {
            "ok":    False,
            "error": (
                "Stable Diffusion är inte tillgänglig. "
                "Installera diffusers: pip install diffusers transformers accelerate"
            ),
        }

    logger.info(
        f"[ImgGen] Generating image  steps={SD_STEPS}  cfg={SD_CFG_SCALE}  "
        f"size={SD_WIDTH}x{SD_HEIGHT}  prompt='{prompt}'"
    )
    try:
        result = pipe(
            prompt,
            negative_prompt     = negative_prompt or SD_NEGATIVE_DEFAULT,
            num_inference_steps = SD_STEPS,
            guidance_scale      = SD_CFG_SCALE,
            width               = SD_WIDTH,
            height              = SD_HEIGHT,
        )
        pil_image = result.images[0]
    except Exception as e:
        logger.error(f"[ImgGen] Pipeline inference failed: {e}")
        return {"ok": False, "error": f"Bildgenerering misslyckades: {e}"}

    buf = io.BytesIO()
    pil_image.save(buf, format="PNG")
    png_bytes = buf.getvalue()
    b64_str   = base64.b64encode(png_bytes).decode()
    img_path  = _save_generated_image(png_bytes)

    logger.info(f"[ImgGen] Done. base64 length={len(b64_str)}")
    return {
        "ok":     True,
        "b64":    b64_str,
        "path":   img_path,
        "prompt": prompt,
    }


# ════════════════════════════════════════════════════════════════════════
# STAGE 1 — Language Detection
# ════════════════════════════════════════════════════════════════════════
SW_WORDS = {
    "och", "att", "det", "är", "en", "ett", "i", "på", "av", "för", "med",
    "som", "har", "inte", "till", "om", "men", "kan", "vi", "de", "så",
    "den", "var", "sig", "eller", "han", "hon", "man", "vad", "nu", "här",
    "när", "hur", "hej", "tack", "ja", "nej", "också", "mycket", "lite",
    "bra", "stor", "liten", "hjälp", "fråga", "svar", "tid", "dag", "år",
    "vill", "ska", "måste", "kanske", "bara", "detta", "dessa", "vilket",
    "vem", "vilka", "varför", "sedan", "efter", "dock", "eftersom", "trots",
    "medan", "alltså", "därför", "vidare", "dessutom", "emellertid",
    "däremot", "således", "exempelvis",
}

NON_SW_PATTERNS = [
    r"\bthe\b", r"\band\b", r"\bwhat\b", r"\bhow\b", r"\bwhy\b",
    r"\bbonjour\b", r"\bmerci\b", r"\bgracias\b", r"\bdanke\b",
    r"\bhello\b", r"\bplease\b", r"\bthank you\b", r"\bthis is\b",
    r"\bthat is\b", r"\bwith the\b", r"\bfrom the\b",
]


def stage1_detect(text: str) -> dict:
    t = text.lower().strip()
    if len(t) < 2:
        return {"is_swedish": True, "confidence": 1.0, "method": "too_short"}

    for pat in NON_SW_PATTERNS:
        if re.search(pat, t, re.IGNORECASE):
            return {"is_swedish": False, "confidence": 0.05, "method": "pattern_reject"}

    words    = re.findall(r"\b\w+\b", t)
    sw_words = sum(1 for w in words if w in SW_WORDS)
    sw_ratio = sw_words / max(len(words), 1)
    sv_chars = sum(1 for c in t if c in "åäöÅÄÖ")
    char_sc  = min(sv_chars / max(len(t), 1) * 15, 1.0)

    kb_score = 0.5
    if _tok and len(t) > 6:
        try:
            tokens   = _tok.tokenize(t[:300])
            n        = max(len(tokens), 1)
            unk_rate = sum(1 for tk in tokens if tk == "[UNK]") / n
            avg_len  = sum(len(tk.replace("##", "")) for tk in tokens) / n
            len_sc   = min(avg_len / 5.0, 1.0)
            kb_score = max(0.0, (1.0 - unk_rate * 2.5) * 0.65 + len_sc * 0.35)
        except Exception:
            pass

    confidence = sw_ratio * 0.30 + char_sc * 0.25 + kb_score * 0.45
    is_swedish = confidence > 0.20 or sw_words >= 2 or sv_chars >= 2

    return {
        "is_swedish": is_swedish,
        "confidence": round(min(confidence, 1.0), 3),
        "method":     "kb_lab_tokenizer",
        "kb_score":   round(kb_score, 3),
        "sw_words":   sw_words,
        "sv_chars":   sv_chars,
    }


# ════════════════════════════════════════════════════════════════════════
# STAGE 2 — Spell-Check & Sentence Repair
# ════════════════════════════════════════════════════════════════════════
SW_VERBS = {
    "är", "var", "blir", "blev", "har", "hade", "kan", "kunde", "vill",
    "ville", "ska", "skulle", "måste", "får", "fick", "gör", "gjorde",
    "säger", "sa", "kommer", "kom", "vet", "visste", "tänker", "tänkte",
    "ser", "såg", "heter", "hette", "finns", "fanns", "verkar", "verkade",
    "behöver", "behövde", "arbetar", "bor", "bodde", "hoppas", "tror",
    "tycker", "tyckte", "förstår", "förstod", "menar", "menade", "önskar",
}

TECH_TERMS = {
    "printf", "async", "await", "usestate", "useeffect", "usereducer",
    "const", "null", "true", "false", "undefined", "api", "http", "https",
    "json", "xml", "html", "css", "sql", "url", "uri", "jwt", "oauth",
    "api", "rest", "graphql", "git", "npm", "pip", "docker", "kubernetes",
    "python", "javascript", "typescript", "java", "golang", "rust", "swift",
    "react", "angular", "vue", "flask", "django", "fastapi", "express",
    "llama", "bert", "gpt", "llm", "ollama", "huggingface",
    "linux", "ubuntu", "debian", "windows", "macos",
    "int", "str", "bool", "float", "list", "dict", "tuple", "set",
    "class", "def", "return", "import", "from", "print", "input",
    "for", "while", "if", "else", "elif", "try", "except", "with", "as",
    "function", "var", "let", "throw", "catch", "new", "this",
    "diffusers", "stablediffusion", "stable", "diffusion", "pipeline",
    "cfg", "steps", "sampler", "lora", "checkpoint", "vae",
}


def stage2_repair(text: str) -> dict:
    words       = text.split()
    repaired    = list(words)
    corrections = []

    if _fill_mask and _tok and len(text) <= 350:
        for i, word in enumerate(words):
            clean = re.sub(r"[^\w]", "", word, flags=re.UNICODE)

            if clean.lower() in TECH_TERMS:
                continue
            if len(clean) <= 2 or clean.isdigit():
                continue
            if re.search(r"[A-Z]", clean[1:]) or "_" in clean:
                continue

            solo_toks = _tok.tokenize(clean)
            is_unk    = ("[UNK]" in solo_toks
                         or all(tk == "[UNK]" or tk.startswith("##")
                                for tk in solo_toks))
            if not is_unk:
                continue

            ctx    = list(words)
            ctx[i] = _tok.mask_token
            masked = " ".join(ctx)

            try:
                preds      = _fill_mask(masked)
                best_tok   = preds[0]["token_str"].strip()
                best_score = preds[0]["score"]

                if (best_score > 0.30
                        and best_tok.lower() != clean.lower()
                        and re.search(r"[a-zA-ZåäöÅÄÖ]", best_tok)
                        and len(best_tok) > 1):
                    suffix      = re.search(r"[^\w]+$", word, re.UNICODE)
                    repaired[i] = best_tok + (suffix.group() if suffix else "")
                    corrections.append({
                        "original":   word,
                        "corrected":  best_tok,
                        "confidence": round(best_score, 3),
                        "reason":     "kb_lab_unk_fill_mask",
                    })
            except Exception:
                continue

    repaired_text = " ".join(repaired)
    was_changed   = repaired_text.strip().lower() != text.strip().lower()
    if was_changed and not repaired_text.rstrip().endswith((".", "?", "!")):
        repaired_text += "?"

    words_lower = {w.lower().rstrip(".,!?;:") for w in repaired}
    has_verb    = bool(words_lower & SW_VERBS)
    is_fragment = not has_verb and len(repaired) <= 6

    return {
        "repaired_text": repaired_text,
        "corrections":   corrections,
        "is_fragment":   is_fragment,
        "was_changed":   was_changed,
    }


# ════════════════════════════════════════════════════════════════════════
# STAGE 3 — NER + Question Understanding
# ════════════════════════════════════════════════════════════════════════
Q_PATTERNS = {
    "definition":     [r"\bvad är\b", r"\bberätta om\b", r"\bförklara\b",
                       r"\bdefiniera\b", r"\binnebär\b", r"\bbetyder\b"],
    "factual":        [r"\bnär\b", r"\bvar\b", r"\bvem\b", r"\bhur många\b",
                       r"\bhur mycket\b", r"\bvilket år\b", r"\bvilken dag\b"],
    "procedural":     [r"\bhur gör man\b", r"\bhur fungerar\b", r"\bhur kan\b",
                       r"\bsteg för steg\b", r"\bhur ska\b"],
    "comparison":     [r"\bjämför\b", r"\bskillnad\b", r"\bbättre\b",
                       r"\bsämre\b", r"\bvs\b"],
    "opinion":        [r"\btycker\b", r"\banser\b", r"\brekommenderar\b",
                       r"\bbästa\b", r"\bsämsta\b"],
    "current_events": [r"\bidag\b", r"\bnyheter?\b", r"\bsenaste\b",
                       r"\bnuvarande\b", r"\bigår\b", r"\bimorgon\b"],
}

ANSWER_STYLE = {
    "definition":        "Ge en tydlig definition med ett konkret exempel.",
    "factual":           "Ge ett precist faktasvar med specifika siffror/datum om möjligt.",
    "procedural":        "Svara steg-för-steg. Numrera varje steg.",
    "comparison":        "Jämför tydligt med för- och nackdelar.",
    "opinion":           "Ge en välmotiverad rekommendation.",
    "current_events":    "Använd den senaste tillgängliga informationen. Ange datum.",
    "image_generation":  "Generera en bild baserat på beskrivningen.",
    "general":           "Svara informativt och tydligt på svenska.",
}


def stage3_understand(text: str) -> dict:
    t = text.lower()

    q_type = "general"
    for qt, patterns in Q_PATTERNS.items():
        if any(re.search(p, t) for p in patterns):
            q_type = qt
            break

    entities = []
    if _ner_pipe:
        try:
            raw  = _ner_pipe(text[:512])
            seen = set()
            for ent in raw:
                key = (ent["entity_group"], ent["word"].lower().strip())
                if key in seen or float(ent["score"]) < 0.65:
                    continue
                seen.add(key)
                entities.append({
                    "type":  ent["entity_group"],
                    "value": ent["word"].strip(),
                    "score": round(float(ent["score"]), 3),
                })
        except Exception as e:
            logger.warning(f"[S3] NER error: {e}")

    if q_type == "general":
        if any(e["type"] == "PER" for e in entities):
            q_type = "factual"
        elif any(e["type"] == "TME" for e in entities):
            q_type = "current_events"

    ctx_parts = []
    for etype, label in [("PER", "Person"), ("ORG", "Organisation"),
                          ("LOC", "Plats"), ("TME", "Tid")]:
        vals = [e["value"] for e in entities if e["type"] == etype]
        if vals:
            ctx_parts.append(f"{label}: {', '.join(vals[:3])}")

    return {
        "entities":      entities,
        "question_type": q_type,
        "answer_style":  ANSWER_STYLE.get(q_type, ANSWER_STYLE["general"]),
        "context_hint":  " | ".join(ctx_parts),
    }


# ════════════════════════════════════════════════════════════════════════
# STAGE 4 — Query-Focused Extractive Summarisation
# ════════════════════════════════════════════════════════════════════════

def _reply_has_structured_content(text: str) -> bool:
    if re.search(r"```|~~~", text):
        return True
    if re.search(r"^(    |\t)\S", text, re.MULTILINE):
        return True
    if re.search(r"^\s*\d+[.)]\s+\S", text, re.MULTILINE):
        return True
    if re.search(r"^\s*[-*•]\s+\S", text, re.MULTILINE):
        return True
    return False


def _mean_pool(model_out, attn_mask):
    tok_embs = model_out.last_hidden_state
    mask     = attn_mask.unsqueeze(-1).float()
    return (tok_embs * mask).sum(dim=1) / mask.sum(dim=1).clamp(min=1e-9)


def _embed(texts: list):
    if not _bert or not _tok or not texts:
        return None
    try:
        enc = _tok(texts, padding=True, truncation=True,
                   max_length=128, return_tensors="pt")
        with torch.no_grad():
            out = _bert(**enc)
        return _mean_pool(out, enc["attention_mask"])
    except Exception as e:
        logger.warning(f"[S4] embed error: {e}")
        return None


def _split_sentences(text: str) -> list:
    parts = re.split(r"(?<=[.!?])\s+(?=[A-ZÅÄÖ\"'(])", text)
    return [p.strip() for p in parts if len(p.strip()) > 18]


def stage4_summarise(llama_reply: str, query: str, max_sents: int = 5) -> dict:
    if _reply_has_structured_content(llama_reply):
        logger.info("[S4] Skipping summarisation — structured content detected")
        return {
            "summary":    llama_reply,
            "used_kb":    False,
            "summarised": False,
            "reason":     "structured_content_preserved",
        }

    sents = _split_sentences(llama_reply)

    if not sents:
        return {"summary": llama_reply, "used_kb": False, "summarised": False}

    if len(sents) <= 2:
        return {
            "summary":    llama_reply,
            "used_kb":    False,
            "summarised": False,
            "reason":     "already_short",
        }

    if len(sents) <= max_sents:
        return {"summary": llama_reply, "used_kb": False, "summarised": False}

    q_emb = _embed([query])
    if q_emb is not None:
        s_embs = _embed(sents)
        if s_embs is not None:
            sims = F.cosine_similarity(
                s_embs,
                q_emb.expand(s_embs.size(0), -1),
                dim=1,
            ).tolist()

            scored = []
            for i, (sent, sim) in enumerate(zip(sents, sims)):
                pos_bonus = max(0.0, 1.0 - i * 0.07)
                len_bonus = min(len(sent) / 120, 1.0)
                score     = sim * 0.70 + pos_bonus * 0.20 + len_bonus * 0.10
                scored.append((i, sent, score))

            top     = sorted(scored, key=lambda x: x[2], reverse=True)[:max_sents]
            top     = sorted(top, key=lambda x: x[0])
            summary = " ".join(s[1] for s in top)
            return {
                "summary":    summary,
                "used_kb":    True,
                "summarised": True,
                "method":     "kb_lab_cosine_similarity",
                "orig":       len(sents),
                "kept":       len(top),
            }

    return {
        "summary":    " ".join(sents[:max_sents]),
        "used_kb":    False,
        "summarised": True,
        "method":     "first_n_fallback",
        "orig":       len(sents),
        "kept":       max_sents,
    }


# ════════════════════════════════════════════════════════════════════════
# Locale
# ════════════════════════════════════════════════════════════════════════
import locale

try:
    locale.setlocale(locale.LC_TIME, "sv_SE.UTF-8")
except Exception:
    try:
        locale.setlocale(locale.LC_TIME, "sv_SE")
    except Exception:
        pass


# ════════════════════════════════════════════════════════════════════════
# Tavily Web Search
# ════════════════════════════════════════════════════════════════════════
TAVILY_URL = "https://api.tavily.com/search"

_MEMORY_RECALL_PATTERNS = [
    r"\bmin\s+(första|förra|senaste|andra|tredje)\s+fråga\b",
    r"\bvad\s+(frågade|sa|sade|nämnde)\s+jag\b",
    r"\bvem\s+är\s+det\s+(vi|jag)\s+pratar\s+om\b",
    r"\bhur\s+många\s+(frågor|meddelanden|svar)\s+(har\s+vi|skickade\s+jag)\b",
    r"\bvad\s+har\s+vi\s+pratat\s+om\b",
    r"\bsammanfatta\s+(vår|denna|den\s+här)\s+(konversation|chatt|diskussion)\b",
    r"\bvad\s+diskuterade\s+vi\b",
    r"\bvad\s+sa\s+du\s+(om|angående|till)\b",
    r"\bvilket\s+ämne\s+(har\s+vi|diskuterar\s+vi)\b",
]

_TAVILY_TRIGGERS = [
    "idag", "igår", "imorgon", "nu", "just nu", "senaste", "nyligen",
    "nyheter", "nyhet", "pris", "kostar", "kostnad",
    "väder", "vädret", "temperatur", "resultat",
    "match", "score", "live", "vann", "vinner", "aktier", "börs",
    "2024", "2025", "2026", "händer",
    "hur mycket kostar", "senaste version", "uppdatering",
    "statsminister", "president", "vd", "regering",
]

DIRECT_ANSWER_TRIGGERS: dict = {
    "time_date": [
        "klockan", "hur dags", "vad är tiden", "tiden nu",
        "dagens datum", "vilket datum", "vad är datumet",
        "vilken dag", "vilken veckodag", "vad är det för dag",
        "tid och datum", "datum och tid", "dag och tid",
        "vad är det för år",
    ],
    "weather": [
        "väder", "vädret", "temperatur", "regnar", "snöar",
        "sol", "molnigt", "vindstyrka", "luftfuktighet",
        "grader", "celsius", "fahrenheit", "prognos", "regn",
        "snö", "storm", "åska", "dimma", "hagel",
    ],
    "news": [
        "nyheter", "nyhet", "senaste", "händer", "händelser",
        "breaking", "aktuellt", "rapporterar", "enligt uppgifter",
    ],
    "prices": [
        "kostar", "pris", "kostnad", "aktier", "börs",
        "kurs", "inflation", "ränta", "rabatt", "erbjudande",
    ],
    "sports": [
        "match", "score", "vann", "förlorade", "resultat",
        "tabell", "liga", "mål", "poäng", "turnering",
        "vm", "em", "champions league", "serie a", "premier league",
    ],
    "current": [
        "just nu", "idag", "igår", "imorgon", "nuvarande",
        "2024", "2025", "2026", "i år", "förra året",
        "den senaste", "nyligen", "häromdagen",
    ],
    "people_roles": [
        "statsminister", "president", "vd", "regering",
        "minister", "chef", "ordförande", "generalsekreterare",
        "kung", "drottning", "påven", "premiärminister",
    ],
    "tech": [
        "senaste version", "uppdatering", "release", "lansering",
        "ny modell", "ny telefon", "ny produkt",
    ],
}


def _is_direct_answer_query(text: str) -> tuple:
    t = text.lower()
    for category, keywords in DIRECT_ANSWER_TRIGGERS.items():
        if any(kw in t for kw in keywords):
            return True, category
    return False, ""


def tavily_needs_search(text: str) -> bool:
    t = text.lower().strip()

    if any(re.search(p, t) for p in _MEMORY_RECALL_PATTERNS):
        return False

    if _is_direct_answer_query(t)[0]:
        return True

    if any(kw in t for kw in _TAVILY_TRIGGERS):
        return True

    current_patterns = [
        r"\bvad händer\b",
        r"\bvad är\b.*\bidag\b",
        r"\bvad är\b.*\bnu\b",
        r"\bhur är\b.*\bidag\b",
        r"\bsenaste om\b",
        r"\baktuellt\b",
        r"\bjust nu\b",
    ]

    if any(re.search(p, t) for p in current_patterns):
        return True

    return False


def tavily_search(query: str) -> list:
    if not TAVILY_KEY:
        return []
    try:
        r = requests.post(TAVILY_URL, json={
            "api_key": TAVILY_KEY, "query": query,
            "search_depth": "basic", "include_answer": True, "max_results": 5,
        }, timeout=12)
        r.raise_for_status()
        data, out = r.json(), []
        if data.get("answer"):
            out.append({"title": "Direkt svar", "url": "", "snippet": data["answer"]})
        for item in data.get("results", [])[:4]:
            out.append({
                "title":   item.get("title", ""),
                "url":     item.get("url", ""),
                "snippet": item.get("content", "")[:300],
            })
        return out
    except Exception as e:
        logger.warning(f"Tavily: {e}")
        return []


def _build_direct_reply(search_results: list, query: str,
                        bypass_category: str) -> tuple:
    combined_snippets = []

    direct = next(
        (r["snippet"] for r in search_results if r["title"] == "Direkt svar"),
        None,
    )
    if direct:
        combined_snippets.append(direct)

    for r in search_results:
        if r["title"] != "Direkt svar" and r.get("snippet"):
            clean_snippet = re.sub(r"#+\s*", "", r["snippet"])
            clean_snippet = re.sub(r"[*_`]", "", clean_snippet)
            clean_snippet = re.sub(r"\s+", " ", clean_snippet).strip()
            if clean_snippet:
                combined_snippets.append(f"{r['title']}: {clean_snippet}")

    full_text = " ".join(combined_snippets)
    summ = stage4_summarise(full_text, query=query, max_sents=5)

    if summ.get("summarised") and summ["summary"].strip():
        final_reply = summ["summary"]
    elif direct:
        final_reply = direct
    elif combined_snippets:
        final_reply = combined_snippets[0]
    else:
        final_reply = "Tyvärr kunde jag inte hitta aktuell information om det."

    sources = [
        r["title"] for r in search_results
        if r["title"] != "Direkt svar" and r.get("title")
    ]
    if sources:
        final_reply += f"\n\n📌 Källa: {', '.join(sources[:2])}"

    logger.info(
        f"[DirectPath] category={bypass_category} "
        f"snippets={len(combined_snippets)} "
        f"summarised={summ.get('summarised')} "
        f"used_kb={summ.get('used_kb')}"
    )
    return final_reply, summ


# ════════════════════════════════════════════════════════════════════════
# Llama via Ollama — standard (buffered) + streaming
# ════════════════════════════════════════════════════════════════════════
HARMFUL_WORDS = [
    "bomba", "sprängämne", "terrorism", "knark", "droger", "skada någon",
    "döda", "attackera", "hacka", "malware", "virus", "ransomware", "självmord",
]


def is_harmful(text: str) -> bool:
    return any(h in text.lower() for h in HARMFUL_WORDS)


def build_system_prompt(understood: dict, is_fragment: bool) -> str:
    lines = [
        "Du är en hjälpsam, vänlig och kunnig AI-assistent som alltid svarar på svenska.",
        "",
        "=== VIKTIGT: AKTUELL INFORMATION ===",
        "För frågor om tid, datum, väder, nyheter eller annan aktuell information:",
        "Svara ALLTID baserat på webbsökningsresultaten som tillhandahålls.",
        "Använd ALDRIG din interna kunskap för aktuell information — den kan vara felaktig.",
        "",
        f"=== FRÅGETYP (KB-Lab) ===",
        f"Typ   : {understood['question_type']}",
        f"Stil  : {understood['answer_style']}",
    ]
    if understood["context_hint"]:
        lines.append(f"Entiteter: {understood['context_hint']}")
    if is_fragment:
        lines.append("OBS: Frågan verkade ofullständig — tolka välvilligt och svara på det troliga syftet.")

    lines += [
        "",
        "=== SVARSREGLER ===",
        "- Svara ALLTID på korrekt, naturlig svenska oavsett frågans språk.",
        "- Håll svaret enkelt, tydligt och lättförståeligt.",
        "- Använd korta meningar. Undvik onödigt tekniskt språk.",
        "- Använd radbrytningar för långa svar. Numrera steg om det är en guide.",
        "- Använd ALLTID markdown-formatering: ```kod``` för kodexempel, **fet** för viktiga begrepp.",
        "- Bevara ALLTID korrekt kodindragning i kodblock.",
        "- Om webbsökningsresultat finns — använd dem och nämn källan kort.",
        "- Vet du inte svaret — säg 'Jag vet tyvärr inte det.' Hitta inte på.",
        "- Vägra ARTIGT skadliga, olagliga eller stötande förfrågningar.",
        "- Avsluta ALDRIG med 'Har du fler frågor?' eller liknande fraser.",
        "",
        "=== KONVERSATIONSMINNE ===",
        "- Du har tillgång till hela konversationshistoriken i den aktuella tråden.",
        "- Använd tidigare meddelanden för att förstå pronomen och implicita referenser.",
        "- Om användaren syftar på något som nämndes tidigare, svara med den kontexten i åtanke.",
    ]
    return "\n".join(lines)


def ollama_generate(history: list, system: str, search_ctx: str = "",
                    num_predict: int = 768,
                    image_b64_list: list | None = None) -> str:
    """Buffered (non-streaming) Ollama call — used for non-streaming /api/chat."""
    if search_ctx:
        system += f"\n\n--- Tavily webbinformation ---\n{search_ctx}\n---"

    # ── Vision path ───────────────────────────────────────────────────────────
    if image_b64_list:
        user_prompt = ""
        for m in reversed(history):
            if m["role"] == "user":
                user_prompt = m["content"]
                break
        if not user_prompt:
            user_prompt = "Beskriv bilden."

        payload = {
            "model":   OLLAMA_MODEL,
            "prompt":  user_prompt,
            "system":  system,
            "images":  image_b64_list,
            "stream":  True,
            "options": {"temperature": 0.7, "num_predict": num_predict},
        }
        full = []
        try:
            with requests.post(
                f"{OLLAMA_BASE}/api/generate",
                json=payload, stream=True,
                timeout=(15, 180),
            ) as resp:
                resp.raise_for_status()
                for raw in resp.iter_lines():
                    if not raw:
                        continue
                    try:
                        chunk = json.loads(raw)
                        tok   = chunk.get("response", "")
                        if tok:
                            full.append(tok)
                        if chunk.get("done"):
                            break
                    except json.JSONDecodeError:
                        continue
            result = "".join(full).strip()
            return result if result else "Kunde inte analysera bilden. Kontrollera att llava är laddad."
        except requests.exceptions.ConnectionError:
            return "Kan inte nå Ollama. Starta med: ollama serve"
        except requests.exceptions.Timeout:
            partial = "".join(full).strip()
            return (partial + "\n_(Svaret avbröts)_") if partial else "Timeout vid bildanalys."
        except Exception as e:
            return f"Fel vid bildanalys: {e}"

    # ── Text path ─────────────────────────────────────────────────────────────
    messages = [{"role": "system", "content": system}] + list(history)
    payload = {
        "model":    OLLAMA_MODEL,
        "messages": messages,
        "stream":   True,
        "options":  {"temperature": 0.7, "num_predict": num_predict, "num_ctx": 8192},
    }
    full = []
    try:
        with requests.post(
            f"{OLLAMA_BASE}/api/chat",
            json=payload, stream=True,
            timeout=(10, 180),
        ) as resp:
            resp.raise_for_status()
            for raw in resp.iter_lines():
                if not raw:
                    continue
                try:
                    chunk = json.loads(raw)
                    tok   = chunk.get("message", {}).get("content", "")
                    if tok:
                        full.append(tok)
                    if chunk.get("done"):
                        break
                except json.JSONDecodeError:
                    continue
        return "".join(full).strip()

    except requests.exceptions.ConnectionError:
        return "Kan inte nå Ollama. Starta med: ollama serve"
    except requests.exceptions.Timeout:
        partial = "".join(full).strip()
        return (partial + "\n_(Svaret avbröts — modellen var för långsam)_") \
               if partial else "Timeout. Prova: ollama pull llama3.2:1b"
    except Exception as e:
        logger.error(f"Ollama: {e}")
        return f"Fel vid Ollama-anrop: {e}"


def ollama_generate_stream(history: list, system: str, search_ctx: str = "",
                           num_predict: int = 768,
                           stop_event: threading.Event | None = None):
    """
    [STR-1] Generator that yields raw token strings from Ollama one-by-one.
    Stops early if stop_event is set.
    """
    if search_ctx:
        system += f"\n\n--- Tavily webbinformation ---\n{search_ctx}\n---"

    messages = [{"role": "system", "content": system}] + list(history)
    payload  = {
        "model":    OLLAMA_MODEL,
        "messages": messages,
        "stream":   True,
        "options":  {"temperature": 0.7, "num_predict": num_predict, "num_ctx": 8192},
    }

    try:
        with requests.post(
            f"{OLLAMA_BASE}/api/chat",
            json=payload, stream=True,
            timeout=(10, 300),
        ) as resp:
            resp.raise_for_status()
            for raw in resp.iter_lines():
                if stop_event and stop_event.is_set():
                    logger.info("[Stream] Stop event received — aborting generation")
                    return

                if not raw:
                    continue
                try:
                    chunk = json.loads(raw)
                    tok   = chunk.get("message", {}).get("content", "")
                    if tok:
                        yield tok
                    if chunk.get("done"):
                        return
                except json.JSONDecodeError:
                    continue

    except requests.exceptions.ConnectionError:
        yield "\n\n⚠️ Kan inte nå Ollama. Starta med: `ollama serve`"
    except requests.exceptions.Timeout:
        yield "\n\n⚠️ _(Svaret avbröts — modellen var för långsam)_"
    except Exception as e:
        yield f"\n\n⚠️ Fel: {e}"


# ════════════════════════════════════════════════════════════════════════
# Ollama availability check
# ════════════════════════════════════════════════════════════════════════

def ollama_available() -> bool:
    try:
        r = requests.get(f"{OLLAMA_BASE}/api/tags", timeout=(2, 4))
        if r.status_code != 200:
            return False
        models = r.json().get("models", [])
        names = [m.get("name", "") for m in models]
        return any(OLLAMA_MODEL in n or n in OLLAMA_MODEL for n in names)
    except Exception:
        return False


_VISION_MODEL_NAMES = {
    "llava", "bakllava", "llava-llama3", "llava-phi3",
    "moondream", "cogvlm", "llama3.2-vision", "minicpm-v",
    "qwen2-vl", "internvl", "phi3-vision",
}


def ollama_model_is_vision() -> bool:
    model_lower = OLLAMA_MODEL.lower()
    return any(v in model_lower for v in _VISION_MODEL_NAMES)


# ════════════════════════════════════════════════════════════════════════
# Thread helpers
# ════════════════════════════════════════════════════════════════════════

def _now() -> str:
    return datetime.utcnow().isoformat() + "Z"


def _make_thread(title: str = "Ny konversation") -> dict:
    return {
        "id":                str(uuid.uuid4()),
        "title":             title,
        "created_at":        _now(),
        "updated_at":        _now(),
        "messages":          [],
        "summary":           None,
        "_last_summary_at":  0,
        "message_count":     0,
    }


def _auto_title(text: str) -> str:
    words = text.split()[:8]
    title = " ".join(words)
    return (title[:52] + "…") if len(title) > 52 else title


def _rolling_summary(thread: dict, system_prompt: str) -> str:
    msgs = thread["messages"]
    if len(msgs) < SUMMARY_TRIGGER:
        return system_prompt

    already   = thread.get("summary")
    new_since = len(msgs) - thread.get("_last_summary_at", 0)
    if already and new_since < 10:
        return (system_prompt
                + f"\n\n=== SAMMANFATTNING AV TIDIGARE KONVERSATION ===\n{already}")

    logger.info(f"[Thread] Rolling summary for thread {thread['id']}")
    history_text = "\n".join(
        f"{'Användare' if m['role'] == 'user' else 'Assistent'}: {m['content']}"
        for m in msgs[:-10]
    )
    sum_prompt = (
        "Sammanfatta följande konversation på svenska i max 5 meningar. "
        "Fokusera på ämnen, beslut och nyckelinformation som är viktig "
        "för att förstå de följande meddelandena:\n\n" + history_text
    )
    summary_reply = ollama_generate(
        [{"role": "user", "content": sum_prompt}],
        "Du är en hjälpsam assistent. Svara alltid på svenska.",
    )

    thread["summary"]          = summary_reply
    thread["_last_summary_at"] = len(msgs)
    thread["messages"]         = msgs[-10:]

    return (system_prompt
            + f"\n\n=== SAMMANFATTNING AV TIDIGARE KONVERSATION ===\n{summary_reply}")


def _build_followup_hint(thread: dict, current_msg: str) -> str:
    msgs = thread["messages"]
    if len(msgs) < 2:
        return ""

    recent = msgs[-6:]
    topics = []
    for m in recent:
        if m["role"] == "assistant":
            first = re.split(r"[.!?\n]", m["content"].strip())[0]
            if first and len(first) > 10:
                topics.append(first[:100])

    if not topics:
        return ""

    return (
        "\n\n=== FÖLJDFRÅGEKONTEXT ===\n"
        "Nedan visas vad som diskuterades nyligen i denna konversation. "
        "Använd detta för att förstå pronomen och implicita referenser:\n"
        + "\n".join(f"- {t}" for t in topics[-3:])
    )


# ════════════════════════════════════════════════════════════════════════
# File upload helpers
# ════════════════════════════════════════════════════════════════════════

def _allowed_file(filename: str) -> bool:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in ALLOWED_EXTS


def _extract_text_from_file(path: str, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext in {"txt", "md", "csv", "json"}:
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()[:6000]
        except Exception as e:
            logger.warning(f"[Upload] text read failed: {e}")
            return ""

    if ext == "pdf" and PDFPLUMBER_OK:
        try:
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages[:10]:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts)[:6000]
        except Exception as e:
            logger.warning(f"[Upload] PDF extract failed: {e}")
            return ""

    if ext == "docx" and DOCX_OK:
        try:
            doc = DocxDocument(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text)[:6000]
        except Exception as e:
            logger.warning(f"[Upload] DOCX extract failed: {e}")
            return ""

    if ext in {"xlsx", "xls"} and OPENPYXL_OK:
        try:
            wb   = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets[:3]:
                for row in sheet.iter_rows(max_row=100, values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append("\t".join(cells))
            return "\n".join(rows)[:6000]
        except Exception as e:
            logger.warning(f"[Upload] XLSX extract failed: {e}")
            return ""

    return ""


def _image_to_base64(path: str):
    try:
        mime, _ = mimetypes.guess_type(path)
        mime     = mime or "image/jpeg"
        with open(path, "rb") as f:
            data = f.read()
        raw_b64 = base64.b64encode(data).decode()
        return f"data:{mime};base64,{raw_b64}", raw_b64
    except Exception as e:
        logger.warning(f"[Upload] base64 encode failed: {e}")
        return None, None


def _build_file_context(saved_files: list) -> tuple:
    text_parts      = []
    image_data_uris = []

    for item in saved_files:
        fname  = item["filename"]
        fpath  = item["path"]
        is_img = item["is_image"]

        if is_img:
            uri, raw_b64 = _image_to_base64(fpath)
            if uri and raw_b64:
                image_data_uris.append({"uri": uri, "b64": raw_b64})
                text_parts.append(f"[Bifogad bild: {fname}]")
        else:
            extracted = _extract_text_from_file(fpath, fname)
            if extracted:
                text_parts.append(
                    f"=== Innehåll i bifogad fil: {fname} ===\n"
                    f"{extracted}\n"
                    f"=== Slut på {fname} ==="
                )
            else:
                text_parts.append(
                    f"[Bifogad fil: {fname} — innehållet kunde inte läsas]"
                )

    return "\n\n".join(text_parts), image_data_uris


def _cleanup_files(saved_files: list) -> None:
    for item in saved_files:
        _safe_unlink(item["path"])


# ════════════════════════════════════════════════════════════════════════
# TTS — KBLab Piper TTS
# ════════════════════════════════════════════════════════════════════════

def _clean_tts_text(text: str, max_chars: int = 900) -> str:
    clean = re.sub(r"```[\s\S]*?```", "", text)
    clean = re.sub(r"`[^`]+`", "", clean)
    clean = re.sub(r"[#*_\[\]()\-~>]", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    return clean[:max_chars]


def synth_tts(text: str) -> str | None:
    clean = _clean_tts_text(text)
    if not clean:
        return None

    if not PIPER_OK or _piper is None:
        logger.warning("TTS: KBLab Piper not available")
        return None

    try:
        wav_path = os.path.join(tempfile.gettempdir(), f"tts_{uuid.uuid4().hex}.wav")
        with wave.open(wav_path, "w") as wf:
            _piper.synthesize(clean, wf)
        logger.info(f"TTS: KBLab Piper  {len(clean)} chars → {wav_path}")
        return wav_path
    except Exception as e:
        logger.warning(f"TTS: KBLab Piper failed: {e}")
        return None


def _wav_to_mp3(wav_path: str) -> str | None:
    try:
        from pydub import AudioSegment
        mp3_path = wav_path.replace(".wav", ".mp3")
        AudioSegment.from_wav(wav_path).export(mp3_path, format="mp3")
        os.unlink(wav_path)
        return mp3_path
    except Exception as e:
        logger.warning(f"WAV→MP3 conversion failed ({e}); serving raw WAV")
        return None


# ════════════════════════════════════════════════════════════════════════
# STT — KBLab wav2vec2-large-voxrex-swedish
# ════════════════════════════════════════════════════════════════════════

TARGET_SR = 16_000


def _convert_to_wav(src_path: str) -> str | None:
    wav_path = src_path + "_converted.wav"

    try:
        from pydub import AudioSegment
        ext = os.path.splitext(src_path)[1].lstrip(".") or "webm"
        seg = AudioSegment.from_file(src_path, format=ext)
        seg = seg.set_frame_rate(TARGET_SR).set_channels(1).set_sample_width(2)
        seg.export(wav_path, format="wav")
        logger.info(f"STT convert: pydub {ext} → wav  ({os.path.getsize(wav_path)} B)")
        return wav_path
    except Exception as e:
        logger.warning(f"STT convert pydub failed: {e}")

    try:
        import subprocess
        result = subprocess.run(
            ["ffmpeg", "-y", "-i", src_path,
             "-ar", str(TARGET_SR), "-ac", "1",
             "-acodec", "pcm_s16le", wav_path],
            capture_output=True, timeout=30,
        )
        if result.returncode == 0 and os.path.exists(wav_path):
            return wav_path
        else:
            logger.warning(f"STT convert ffmpeg stderr: {result.stderr.decode()[:300]}")
    except FileNotFoundError:
        logger.warning("STT convert: ffmpeg not found in PATH")
    except Exception as e:
        logger.warning(f"STT convert ffmpeg subprocess failed: {e}")

    try:
        wf, sr = torchaudio.load(src_path, backend="ffmpeg")
        if sr != TARGET_SR:
            wf = torchaudio.transforms.Resample(sr, TARGET_SR)(wf)
        if wf.shape[0] > 1:
            wf = wf.mean(dim=0, keepdim=True)
        torchaudio.save(wav_path, wf, TARGET_SR)
        return wav_path
    except Exception as e:
        logger.warning(f"STT convert torchaudio ffmpeg backend failed: {e}")

    return None


def _load_wav_safe(wav_path: str):
    import numpy as np

    try:
        import soundfile as sf
        data, sr = sf.read(wav_path, dtype="float32", always_2d=False)
        if data.ndim == 2:
            data = data.mean(axis=1)
        waveform = torch.from_numpy(data).unsqueeze(0)
        return waveform, sr
    except ImportError:
        logger.warning("STT load: soundfile not installed, trying scipy")
    except Exception as e:
        logger.warning(f"STT load: soundfile failed ({e}), trying scipy")

    try:
        from scipy.io import wavfile
        sr, data = wavfile.read(wav_path)
        if data.dtype != np.float32:
            data = data.astype(np.float32) / np.iinfo(data.dtype).max
        if data.ndim == 2:
            data = data.mean(axis=1)
        waveform = torch.from_numpy(data).unsqueeze(0)
        return waveform, sr
    except Exception as e:
        raise RuntimeError(f"All WAV loaders failed. Last error: {e}")


def kb_transcribe(audio_path: str) -> str:
    if not WAV2VEC_OK or _stt_proc is None or _stt_model is None:
        logger.warning("STT: KBLab wav2vec2 model not loaded")
        return ""

    logger.info(f"STT: received file {audio_path}  size={os.path.getsize(audio_path)} B")

    wav_path   = None
    need_clean = False

    if audio_path.lower().endswith(".wav"):
        try:
            wf_test, _ = torchaudio.load(audio_path)
            wav_path   = audio_path
        except Exception:
            pass

    if wav_path is None:
        wav_path   = _convert_to_wav(audio_path)
        need_clean = True
        if wav_path is None:
            return ""

    try:
        waveform, sr = _load_wav_safe(wav_path)
    except Exception as e:
        logger.error(f"STT: WAV load failed: {e}")
        return ""
    finally:
        if need_clean and wav_path and wav_path != audio_path:
            _safe_unlink(wav_path)

    if sr != TARGET_SR:
        try:
            from scipy.signal import resample_poly
            import math
            g = math.gcd(int(sr), TARGET_SR)
            up, down = TARGET_SR // g, int(sr) // g
            arr = waveform.squeeze().numpy()
            arr = resample_poly(arr, up, down).astype("float32")
            waveform = torch.from_numpy(arr).unsqueeze(0)
        except Exception as e:
            logger.warning(f"STT resample scipy failed ({e}), trying torchaudio")
            waveform = torchaudio.transforms.Resample(sr, TARGET_SR)(waveform)

    if waveform.shape[0] > 1:
        waveform = waveform.mean(dim=0, keepdim=True)

    speech = waveform.squeeze().numpy()
    peak   = float(abs(speech).max())
    dur_s  = len(speech) / TARGET_SR

    if dur_s < 0.3:
        return ""
    if peak < 0.0005:
        return ""
    if peak < 0.1:
        speech = speech / (peak + 1e-9)

    try:
        import numpy as np
        inputs = _stt_proc(
            speech,
            sampling_rate=TARGET_SR,
            return_tensors="pt",
            padding=True,
        )
        with torch.no_grad():
            logits = _stt_model(
                inputs.input_values,
                attention_mask=inputs.get("attention_mask"),
            ).logits

        predicted_ids = torch.argmax(logits, dim=-1)
        transcript    = _stt_proc.batch_decode(predicted_ids)[0].strip()
        logger.info(f"STT: transcript='{transcript}'")
        return transcript

    except Exception as e:
        logger.error(f"STT: wav2vec2 inference failed: {e}")
        return ""


def _safe_unlink(path: str) -> None:
    try:
        os.unlink(path)
    except OSError:
        pass


# ── Startup ───────────────────────────────────────────────────────────────────
_persist_load()
load_kb_models()
logger.info(
    f"[ImgGen] diffusers={'OK' if DIFFUSERS_OK else 'NOT INSTALLED'}  "
    f"model={SD_MODEL_ID}  "
    f"device={'cuda' if (TRANSFORMERS_OK and torch.cuda.is_available()) else 'cpu'}"
)


# ════════════════════════════════════════════════════════════════════════
# /api/chat — Full KB-Lab pipeline (buffered, non-streaming)
# ════════════════════════════════════════════════════════════════════════
def _get_user_id() -> str:
    uid = request.headers.get("X-User-ID", "").strip()
    if uid:
        return uid
    try:
        body = request.get_json(force=True, silent=True) or {}
        uid  = body.get("user_id", "").strip()
        if uid:
            return uid
    except Exception:
        pass
    uid = session.get("session_id", "").strip()
    if uid:
        return uid
    uid = str(uuid.uuid4())
    session["session_id"] = uid
    return uid


@app.route("/api/chat", methods=["POST"])
def chat():

    # ── Parse request ─────────────────────────────────────────────────────────
    if request.content_type and "multipart/form-data" in request.content_type:
        msg       = request.form.get("message", "").strip()
        thread_id = request.form.get("thread_id")
        uploaded  = request.files.getlist("files")
    else:
        data      = request.get_json(force=True)
        msg       = data.get("message", "").strip()
        thread_id = data.get("thread_id")
        uploaded  = []

    sid = _get_user_id()

    if not msg and not uploaded:
        return jsonify({"error": "Tomt meddelande"}), 400

    # ── Save uploaded files ───────────────────────────────────────────────────
    saved_files  = []
    file_summary = []

    for f in uploaded[:MAX_FILES]:
        if not f or not f.filename:
            continue
        if not _allowed_file(f.filename):
            continue
        safe_name = secure_filename(f.filename)
        dest_path = os.path.join(UPLOAD_FOLDER, f"{uuid.uuid4().hex}_{safe_name}")
        try:
            f.seek(0, 2); size = f.tell(); f.seek(0)
            if size > MAX_FILE_BYTES:
                continue
            f.save(dest_path)
        except Exception as e:
            logger.warning(f"[Upload] Save failed for {safe_name}: {e}")
            continue
        ext      = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
        is_image = ext in ALLOWED_IMAGE_EXTS
        saved_files.append({"path": dest_path, "filename": safe_name, "is_image": is_image})
        file_summary.append(safe_name)

    if not msg and file_summary:
        all_images = all(f["is_image"] for f in saved_files)
        if all_images:
            msg = "Beskriv vad du ser i bilden." if len(saved_files) == 1 else "Beskriv vad du ser i bilderna."
        else:
            msg = f"Beskriv och analysera det bifogade innehållet: {', '.join(file_summary)}"

    pinfo = {}

    # ── [IMG-FIX-2] Check for image generation BEFORE language detection ──────
    # This prevents Tavily from intercepting image generation requests.
    # We still enforce Swedish-only AFTER confirming it's an image request.
    if is_image_generation_request(msg) and not saved_files:
        # Now check language — must be Swedish
        lang = stage1_detect(msg)
        pinfo["stage1"] = lang
        logger.info(f"[ImgGen] Image request detected. Swedish={lang['is_swedish']}")

        if not lang["is_swedish"]:
            return jsonify({
                "reply":         "Förlåt, jag förstår bara svenska. Vänligen skriv på svenska.",
                "lang_ok":       False,
                "confidence":    lang["confidence"],
                "pipeline_info": pinfo,
            })

        thread = _resolve_thread(sid, thread_id)
        if thread["message_count"] == 0:
            thread["title"] = _auto_title(msg)

        spell      = stage2_repair(msg)
        clean      = spell["repaired_text"]
        understood = stage3_understand(clean)
        pinfo["stage2"] = {"was_changed": spell["was_changed"], "corrections": spell["corrections"]}
        pinfo["stage3"] = {"question_type": "image_generation"}

        img_prompt     = extract_image_prompt(clean)
        english_prompt = _translate_prompt_to_english(img_prompt)
        logger.info(f"[ImgGen] Prompt SV='{img_prompt}' EN='{english_prompt}'")
        img_result     = generate_image_diffusers(english_prompt)

        if img_result["ok"]:
            data_uri     = f"data:image/png;base64,{img_result['b64']}"
            imggen_reply = (
                f"Här är din genererade bild baserad på: *{img_prompt}*\n\n"
                f"![Genererad bild]({data_uri})"
            )
        else:
            imggen_reply = (
                f"Tyvärr kunde jag inte generera bilden.\n\n"
                f"**Fel:** {img_result.get('error', 'Okänt fel')}\n\n"
                f"Kontrollera att diffusers är installerat:\n"
                f"```\npip install diffusers transformers accelerate\n```"
            )

        thread["messages"].append({"role": "user",      "content": clean})
        thread["messages"].append({"role": "assistant", "content": imggen_reply})
        if len(thread["messages"]) > MAX_THREAD_MESSAGES:
            thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
        thread["message_count"] += 2
        thread["updated_at"]     = _now()
        pinfo["imggen"] = {
            "ok":        img_result["ok"],
            "prompt_sv": img_prompt,
            "prompt_en": english_prompt,
        }
        _persist_save()

        return jsonify({
            "reply":             imggen_reply,
            "lang_ok":           True,
            "search_used":       False,
            "search_results":    [],
            "confidence":        lang["confidence"],
            "entities":          understood["entities"],
            "question_type":     "image_generation",
            "spell_corrections": spell["corrections"],
            "query_repaired":    spell["was_changed"],
            "repaired_query":    clean if spell["was_changed"] else None,
            "is_fragment":       spell["is_fragment"],
            "was_summarised":    False,
            "pipeline_info":     pinfo,
            "thread_id":         thread["id"],
            "thread_title":      thread["title"],
            "image_generated":   img_result["ok"],
            "image_b64":         img_result.get("b64"),
            "image_prompt":      img_prompt,
            "image_prompt_en":   english_prompt,
        })

    # ── Stage 1 — Language Detection (normal path) ────────────────────────────
    lang = stage1_detect(msg)
    pinfo["stage1"] = lang
    logger.info(f"[S1] swedish={lang['is_swedish']} conf={lang['confidence']}")

    if not lang["is_swedish"]:
        _cleanup_files(saved_files)
        return jsonify({
            "reply":         "Förlåt, jag förstår bara svenska. Vänligen skriv på svenska.",
            "lang_ok":       False,
            "confidence":    lang["confidence"],
            "pipeline_info": pinfo,
        })

    if is_harmful(msg):
        _cleanup_files(saved_files)
        return jsonify({"reply": "Tyvärr kan jag inte hjälpa med det.", "lang_ok": True, "refused": True})

    thread = _resolve_thread(sid, thread_id)
    if thread["message_count"] == 0:
        thread["title"] = _auto_title(msg)

    # Stage 2 — Spell-Check
    spell = stage2_repair(msg)
    pinfo["stage2"] = {
        "was_changed": spell["was_changed"],
        "is_fragment": spell["is_fragment"],
        "corrections": spell["corrections"],
    }
    clean = spell["repaired_text"]

    # Stage 3 — NER + Intent
    understood = stage3_understand(clean)
    pinfo["stage3"] = {
        "question_type": understood["question_type"],
        "entity_count":  len(understood["entities"]),
        "context_hint":  understood["context_hint"],
    }

    file_context_text, image_data_uris = ("", [])
    if saved_files:
        file_context_text, image_data_uris = _build_file_context(saved_files)

    # Memory recall
    _memory_patterns = [
        r"\bmin\s+(första|förra|senaste)\s+fråga\b",
        r"\bvad\s+(frågade|sa|sade|nämnde)\s+jag\b",
        r"\bvad\s+har\s+vi\s+pratat\s+om\b",
        r"\bsammanfatta\s+(vår|denna|den\s+här)\s+(konversation|chatt)",
        r"\bhur\s+många\s+(frågor|meddelanden)\s+har\s+vi\b",
    ]
    _clean_lower      = clean.lower()
    _is_memory_recall = any(re.search(p, _clean_lower) for p in _memory_patterns)

    if _is_memory_recall:
        user_msgs = [m["content"] for m in thread["messages"] if m["role"] == "user"]
        if user_msgs:
            if re.search(r"\bmin\s+(första|förra)\s+fråga\b", _clean_lower):
                memory_answer = f'Din första fråga var: "{user_msgs[0]}"'
            elif re.search(r"\bmin\s+senaste\s+fråga\b", _clean_lower):
                memory_answer = f'Din senaste fråga var: "{user_msgs[-1]}"'
            elif re.search(r"\bhur\s+många\s+(frågor|meddelanden)", _clean_lower):
                memory_answer = f"Vi har haft {len(user_msgs)} frågor i den här konversationen."
            elif re.search(r"\bvad\s+har\s+vi\s+pratat\s+om\b|\bsammanfatta", _clean_lower):
                topics = ", ".join(f'"{m}"' for m in user_msgs[:5])
                memory_answer = f"Vi har pratat om följande: {topics}."
            else:
                memory_answer = f'Din första fråga var: "{user_msgs[0]}"'
        else:
            memory_answer = "Du har inte ställt någon fråga ännu."

        thread["messages"].append({"role": "user",      "content": clean})
        thread["messages"].append({"role": "assistant", "content": memory_answer})
        thread["message_count"] += 2
        thread["updated_at"]     = _now()
        _persist_save()
        _cleanup_files(saved_files)
        return jsonify({
            "reply": memory_answer, "lang_ok": True,
            "search_used": False, "search_results": [],
            "confidence": lang["confidence"], "entities": [],
            "question_type": "factual", "spell_corrections": spell["corrections"],
            "query_repaired": spell["was_changed"],
            "repaired_query": clean if spell["was_changed"] else None,
            "is_fragment": False, "was_summarised": False,
            "pipeline_info": pinfo, "thread_id": thread["id"],
            "thread_title": thread["title"],
        })

    # Tavily web search
    should_bypass, bypass_category = _is_direct_answer_query(clean)
    search_ctx, search_results, search_used = "", [], False
    if not saved_files and tavily_needs_search(clean):
        search_results = tavily_search(clean)
        search_used    = bool(search_results)
        if search_results:
            search_ctx = "\n".join(
                f"[{r['title']}]: {r['snippet']}"
                for r in search_results if r.get("snippet")
            )

    # Direct path (Tavily only)
    if should_bypass and search_used and search_results and not saved_files:
        logger.info(f"[DirectPath] ⚡ Bypassing Ollama — category={bypass_category}")
        final_reply, summ = _build_direct_reply(search_results, clean, bypass_category)

        thread["messages"].append({"role": "user",      "content": clean})
        thread["messages"].append({"role": "assistant",  "content": final_reply})
        if len(thread["messages"]) > MAX_THREAD_MESSAGES:
            thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
        thread["message_count"] += 2
        thread["updated_at"]     = _now()
        pinfo["direct_path"] = {
            "bypassed_ollama": True, "category": bypass_category,
            "used_kb": summ.get("used_kb", False), "summarised": summ.get("summarised", False),
        }
        _persist_save()

        return jsonify({
            "reply": final_reply, "lang_ok": True,
            "search_used": True, "search_results": search_results[:3],
            "confidence": lang["confidence"], "entities": understood["entities"],
            "question_type": understood["question_type"],
            "spell_corrections": spell["corrections"],
            "query_repaired": spell["was_changed"],
            "repaired_query": clean if spell["was_changed"] else None,
            "is_fragment": spell["is_fragment"],
            "was_summarised": summ.get("summarised", False),
            "pipeline_info": pinfo, "thread_id": thread["id"],
            "thread_title": thread["title"],
            "direct_path": True, "bypass_category": bypass_category,
        })

    # Ollama availability
    _ollama_up = ollama_available()
    if not _ollama_up:
        logger.warning("[Chat] Ollama unavailable — falling back to Tavily → KB-Lab")
        if not search_used:
            search_results = tavily_search(clean)
            search_used    = bool(search_results)
        if search_used and search_results:
            final_reply, summ = _build_direct_reply(search_results, clean, bypass_category or "fallback")
        else:
            final_reply = (
                "Ollama är inte tillgängligt just nu och jag kunde inte hitta "
                "information via webben heller. Starta Ollama med: ollama serve"
            )
            summ = {"summarised": False, "used_kb": False, "method": "offline_fallback"}

        thread["messages"].append({"role": "user",      "content": clean})
        thread["messages"].append({"role": "assistant",  "content": final_reply})
        if len(thread["messages"]) > MAX_THREAD_MESSAGES:
            thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
        thread["message_count"] += 2
        thread["updated_at"]     = _now()
        _persist_save()
        _cleanup_files(saved_files)

        return jsonify({
            "reply": final_reply, "lang_ok": True,
            "search_used": search_used, "search_results": search_results[:3],
            "confidence": lang["confidence"], "entities": understood["entities"],
            "question_type": understood["question_type"],
            "spell_corrections": spell["corrections"],
            "query_repaired": spell["was_changed"],
            "repaired_query": clean if spell["was_changed"] else None,
            "is_fragment": spell["is_fragment"],
            "was_summarised": summ.get("summarised", False),
            "pipeline_info": pinfo, "thread_id": thread["id"],
            "thread_title": thread["title"],
            "direct_path": True, "bypass_category": "ollama_fallback",
            "ollama_available": False,
        })

    # Normal Ollama path
    base_system   = build_system_prompt(understood, spell["is_fragment"])
    followup_hint = _build_followup_hint(thread, clean)

    if file_context_text:
        base_system += (
            "\n\n=== BIFOGADE FILER ===\n"
            "Användaren har bifogat följande filer. "
            "Använd innehållet nedan för att svara på frågan:\n\n"
            + file_context_text
            + "\n=== SLUT PÅ BIFOGADE FILER ==="
        )

    system = _rolling_summary(thread, base_system + followup_hint)

    history_for_ollama = list(thread["messages"]) + [{"role": "user", "content": clean}]
    if len(history_for_ollama) > MAX_THREAD_MESSAGES:
        history_for_ollama = history_for_ollama[-MAX_THREAD_MESSAGES:]

    raw_b64_images  = [img["b64"] for img in image_data_uris if img.get("b64")]
    _is_vision      = ollama_model_is_vision()
    _vision_warning = ""

    if raw_b64_images:
        if _is_vision:
            system += (
                f"\n\n[{len(raw_b64_images)} bild(er) bifogade. "
                "Beskriv och analysera bilderna noggrant på svenska.]"
            )
        else:
            _vision_warning = (
                f"⚠️ **OBS:** Den aktuella modellen (`{OLLAMA_MODEL}`) stöder inte bildanalys. "
                f"Starta en vision-modell för att analysera bilder, t.ex.:\n"
                f"```\nollama pull llava\nOLLAMA_MODEL=llava python swedishchat.py\n```"
            )

    _question_type  = understood.get("question_type", "factual")
    _msg_len        = len(clean.split())
    if _question_type in ("factual", "definition") and _msg_len <= 10:
        _ollama_predict = 512
    elif _question_type in ("procedural", "comparison"):
        _ollama_predict = 1024
    else:
        _ollama_predict = 768
    if saved_files:
        _ollama_predict = max(_ollama_predict, 1024)

    if _vision_warning and not file_context_text and not msg.strip():
        _cleanup_files(saved_files)
        return jsonify({
            "reply": _vision_warning, "lang_ok": True,
            "search_used": False, "search_results": [],
            "confidence": lang["confidence"], "entities": [],
            "question_type": "general", "spell_corrections": [],
            "query_repaired": False, "repaired_query": None,
            "is_fragment": False, "was_summarised": False,
            "pipeline_info": pinfo, "thread_id": thread["id"],
            "thread_title": thread["title"], "vision_warning": True,
        })

    raw_reply = ollama_generate(
        history_for_ollama, system, search_ctx,
        num_predict=_ollama_predict,
        image_b64_list=raw_b64_images if (raw_b64_images and _is_vision) else None,
    )

    if _vision_warning:
        raw_reply = _vision_warning + "\n\n---\n\n" + raw_reply

    _ollama_failed = (
        "Timeout" in raw_reply or
        "Kan inte nå Ollama" in raw_reply or
        "Fel vid Ollama" in raw_reply
    )
    if _ollama_failed:
        if not search_used:
            search_results = tavily_search(clean)
            search_used    = bool(search_results)
        if search_used and search_results:
            raw_reply, summ = _build_direct_reply(search_results, clean, "ollama_timeout_fallback")
        else:
            summ = {"summarised": False, "used_kb": False, "method": "partial_ollama"}
    else:
        summ = stage4_summarise(raw_reply, query=clean, max_sents=5)

    thread["messages"].append({"role": "user",      "content": clean})
    thread["messages"].append({"role": "assistant", "content": raw_reply})
    if len(thread["messages"]) > MAX_THREAD_MESSAGES:
        thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]

    pinfo["stage4"] = {
        "used_kb": summ.get("used_kb", False), "summarised": summ.get("summarised", False),
        "method": summ.get("method", "none"), "orig": summ.get("orig", 0), "kept": summ.get("kept", 0),
    }
    if saved_files:
        pinfo["files"] = {
            "count": len(saved_files), "names": [f["filename"] for f in saved_files],
            "images": len(raw_b64_images),
        }

    thread["message_count"] += 2
    thread["updated_at"]     = _now()
    _persist_save()
    _cleanup_files(saved_files)

    return jsonify({
        "reply":             summ.get("summary", raw_reply),
        "lang_ok":           True,
        "search_used":       search_used,
        "search_results":    search_results[:3],
        "confidence":        lang["confidence"],
        "entities":          understood["entities"],
        "question_type":     understood["question_type"],
        "spell_corrections": spell["corrections"],
        "query_repaired":    spell["was_changed"],
        "repaired_query":    clean if spell["was_changed"] else None,
        "is_fragment":       spell["is_fragment"],
        "was_summarised":    summ.get("summarised", False),
        "pipeline_info":     pinfo,
        "thread_id":         thread["id"],
        "thread_title":      thread["title"],
        "direct_path":       _ollama_failed,
        "bypass_category":   "ollama_timeout_fallback" if _ollama_failed else None,
        "ollama_available":  True,
        "files_received":    len(saved_files),
    })


# ════════════════════════════════════════════════════════════════════════
# [STR-1] /api/chat/stream — SSE streaming endpoint
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/chat/stream", methods=["POST"])
def chat_stream():
    data      = request.get_json(force=True) or {}
    msg       = data.get("message", "").strip()
    thread_id = data.get("thread_id")
    sid       = _get_user_id()

    if not msg:
        return jsonify({"error": "Tomt meddelande"}), 400

    def _sse(obj: dict) -> str:
        return f"data: {json.dumps(obj, ensure_ascii=False)}\n\n"

    def generate():
        # ── [IMG-FIX-2] Image generation check BEFORE language check ──────────
        if is_image_generation_request(msg):
            lang = stage1_detect(msg)
            if not lang["is_swedish"]:
                yield _sse({"type": "error", "text": "Förlåt, jag förstår bara svenska. Vänligen skriv på svenska."})
                return

            thread = _resolve_thread(sid, thread_id)
            if thread["message_count"] == 0:
                thread["title"] = _auto_title(msg)

            yield _sse({"type": "meta", "thread_id": thread["id"], "thread_title": thread["title"]})

            spell      = stage2_repair(msg)
            clean      = spell["repaired_text"]
            understood = stage3_understand(clean)

            img_prompt     = extract_image_prompt(clean)
            english_prompt = _translate_prompt_to_english(img_prompt)
            logger.info(f"[ImgGen/Stream] Prompt SV='{img_prompt}' EN='{english_prompt}'")

            yield _sse({"type": "token", "text": f"Genererar bild för: *{img_prompt}*\n\n"})

            img_result = generate_image_diffusers(english_prompt)

            if img_result["ok"]:
                data_uri  = f"data:image/png;base64,{img_result['b64']}"
                full_reply = (
                    f"Här är din genererade bild baserad på: *{img_prompt}*\n\n"
                    f"![Genererad bild]({data_uri})"
                )
            else:
                full_reply = (
                    f"Tyvärr kunde jag inte generera bilden.\n\n"
                    f"**Fel:** {img_result.get('error', 'Okänt fel')}"
                )

            yield _sse({"type": "token", "text": full_reply})

            thread["messages"].append({"role": "user",      "content": clean})
            thread["messages"].append({"role": "assistant", "content": full_reply})
            if len(thread["messages"]) > MAX_THREAD_MESSAGES:
                thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
            thread["message_count"] += 2
            thread["updated_at"]     = _now()
            _persist_save()

            yield _sse({
                "type":          "done",
                "thread_id":     thread["id"],
                "thread_title":  thread["title"],
                "question_type": "image_generation",
                "search_used":   False,
                "was_stopped":   False,
                "image_generated": img_result["ok"],
            })
            return

        # ── Stage 1 — Language ────────────────────────────────────────────────
        lang = stage1_detect(msg)
        if not lang["is_swedish"]:
            yield _sse({"type": "error", "text": "Förlåt, jag förstår bara svenska. Vänligen skriv på svenska."})
            return

        if is_harmful(msg):
            yield _sse({"type": "error", "text": "Tyvärr kan jag inte hjälpa med det."})
            return

        thread = _resolve_thread(sid, thread_id)
        if thread["message_count"] == 0:
            thread["title"] = _auto_title(msg)

        yield _sse({"type": "meta", "thread_id": thread["id"], "thread_title": thread["title"]})

        # ── Stage 2 — Spell-check ─────────────────────────────────────────────
        spell      = stage2_repair(msg)
        clean      = spell["repaired_text"]

        # ── Stage 3 — NER + Intent ────────────────────────────────────────────
        understood = stage3_understand(clean)

        # ── Memory recall short-circuit ───────────────────────────────────────
        _memory_patterns = [
            r"\bmin\s+(första|förra|senaste)\s+fråga\b",
            r"\bvad\s+(frågade|sa|sade|nämnde)\s+jag\b",
            r"\bvad\s+har\s+vi\s+pratat\s+om\b",
            r"\bsammanfatta\s+(vår|denna|den\s+här)\s+(konversation|chatt)",
            r"\bhur\s+många\s+(frågor|meddelanden)\s+har\s+vi\b",
        ]
        _clean_lower = clean.lower()
        _is_memory_recall = any(re.search(p, _clean_lower) for p in _memory_patterns)

        if _is_memory_recall:
            user_msgs = [m["content"] for m in thread["messages"] if m["role"] == "user"]
            if user_msgs:
                if re.search(r"\bmin\s+(första|förra)\s+fråga\b", _clean_lower):
                    answer = f'Din första fråga var: "{user_msgs[0]}"'
                elif re.search(r"\bmin\s+senaste\s+fråga\b", _clean_lower):
                    answer = f'Din senaste fråga var: "{user_msgs[-1]}"'
                elif re.search(r"\bhur\s+många", _clean_lower):
                    answer = f"Vi har haft {len(user_msgs)} frågor i den här konversationen."
                else:
                    topics = ", ".join(f'"{m}"' for m in user_msgs[:5])
                    answer = f"Vi har pratat om följande: {topics}."
            else:
                answer = "Du har inte ställt någon fråga ännu."

            for word in answer.split():
                yield _sse({"type": "token", "text": word + " "})

            thread["messages"].append({"role": "user",      "content": clean})
            thread["messages"].append({"role": "assistant", "content": answer})
            thread["message_count"] += 2
            thread["updated_at"]     = _now()
            _persist_save()
            yield _sse({"type": "done", "thread_id": thread["id"], "question_type": "factual"})
            return

        # ── Tavily search ─────────────────────────────────────────────────────
        should_bypass, bypass_category = _is_direct_answer_query(clean)
        search_ctx, search_results, search_used = "", [], False

        if tavily_needs_search(clean):
            search_results = tavily_search(clean)
            search_used    = bool(search_results)
            if search_results:
                search_ctx = "\n".join(
                    f"[{r['title']}]: {r['snippet']}"
                    for r in search_results if r.get("snippet")
                )

        # ── Direct path ───────────────────────────────────────────────────────
        if should_bypass and search_used and search_results:
            final_reply, _ = _build_direct_reply(search_results, clean, bypass_category)
            for word in final_reply.split():
                yield _sse({"type": "token", "text": word + " "})

            thread["messages"].append({"role": "user",      "content": clean})
            thread["messages"].append({"role": "assistant", "content": final_reply})
            if len(thread["messages"]) > MAX_THREAD_MESSAGES:
                thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
            thread["message_count"] += 2
            thread["updated_at"]     = _now()
            _persist_save()
            yield _sse({
                "type": "done", "thread_id": thread["id"],
                "question_type": understood["question_type"],
                "search_used": True,
            })
            return

        # ── Reset stop event for this session ─────────────────────────────────
        stop_evt = _reset_stop_event(sid)

        # ── Build system prompt & history ─────────────────────────────────────
        base_system   = build_system_prompt(understood, spell["is_fragment"])
        followup_hint = _build_followup_hint(thread, clean)
        system        = _rolling_summary(thread, base_system + followup_hint)

        history_for_ollama = list(thread["messages"]) + [{"role": "user", "content": clean}]
        if len(history_for_ollama) > MAX_THREAD_MESSAGES:
            history_for_ollama = history_for_ollama[-MAX_THREAD_MESSAGES:]

        _question_type = understood.get("question_type", "factual")
        _msg_len       = len(clean.split())
        if _question_type in ("factual", "definition") and _msg_len <= 10:
            num_predict = 512
        elif _question_type in ("procedural", "comparison"):
            num_predict = 1024
        else:
            num_predict = 768

        # ── Stream tokens from Ollama ─────────────────────────────────────────
        full_reply_parts = []
        try:
            for tok in ollama_generate_stream(
                history_for_ollama, system, search_ctx,
                num_predict=num_predict,
                stop_event=stop_evt,
            ):
                if stop_evt.is_set():
                    break
                full_reply_parts.append(tok)
                yield _sse({"type": "token", "text": tok})
        except GeneratorExit:
            stop_evt.set()

        full_reply = "".join(full_reply_parts).strip()
        was_stopped = stop_evt.is_set()

        if full_reply:
            thread["messages"].append({"role": "user",      "content": clean})
            thread["messages"].append({"role": "assistant", "content": full_reply})
            if len(thread["messages"]) > MAX_THREAD_MESSAGES:
                thread["messages"] = thread["messages"][-MAX_THREAD_MESSAGES:]
            thread["message_count"] += 2
            thread["updated_at"]     = _now()
            _persist_save()

        yield _sse({
            "type":          "done",
            "thread_id":     thread["id"],
            "thread_title":  thread["title"],
            "question_type": understood["question_type"],
            "search_used":   search_used,
            "was_stopped":   was_stopped,
            "spell_corrections": spell["corrections"],
        })

    return Response(
        stream_with_context(generate()),
        mimetype="text/event-stream",
        headers={
            "Cache-Control":    "no-cache",
            "X-Accel-Buffering": "no",
            "Connection":       "keep-alive",
        },
    )


# ════════════════════════════════════════════════════════════════════════
# [STR-2] /api/chat/stop
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/chat/stop", methods=["POST"])
def chat_stop():
    sid = _get_user_id()
    with _active_gen_lock:
        evt = _active_generations.get(sid)
        if evt:
            evt.set()
            logger.info(f"[Stop] Stop event set for sid={sid[:8]}")
            return jsonify({"status": "stopped"})
    return jsonify({"status": "no_active_generation"})


# ════════════════════════════════════════════════════════════════════════
# [STR-3] /api/threads/<id>/messages/<idx>/edit
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/threads/<thread_id>/messages/<int:msg_index>/edit", methods=["POST"])
def api_edit_message(thread_id: str, msg_index: int):
    thread = ALL_THREADS.get(thread_id)
    if not thread:
        return jsonify({"error": "Tråden hittades inte"}), 404

    msgs = thread["messages"]
    if msg_index < 0 or msg_index >= len(msgs):
        return jsonify({"error": f"Meddelande-index {msg_index} är utanför intervallet"}), 400

    target = msgs[msg_index]
    if target.get("role") != "user":
        return jsonify({"error": "Kan bara redigera användarmeddelanden"}), 400

    body        = request.get_json(force=True) or {}
    new_content = body.get("content", "").strip()
    if not new_content:
        return jsonify({"error": "Innehållet får inte vara tomt"}), 400

    thread["messages"] = msgs[:msg_index]
    thread["message_count"] = len(thread["messages"])
    thread["updated_at"]    = _now()
    _persist_save()

    return jsonify({
        "status":         "edited",
        "thread_id":      thread_id,
        "message_count":  thread["message_count"],
        "edited_content": new_content,
        "next_action":    "resend_to_stream",
    })


# ════════════════════════════════════════════════════════════════════════
# [STR-4] /api/threads/<id>/regenerate
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/threads/<thread_id>/regenerate", methods=["POST"])
def api_regenerate(thread_id: str):
    thread = ALL_THREADS.get(thread_id)
    if not thread:
        return jsonify({"error": "Tråden hittades inte"}), 404

    msgs = thread["messages"]
    last_user_idx = None
    for i in range(len(msgs) - 1, -1, -1):
        if msgs[i]["role"] == "user":
            last_user_idx = i
            break

    if last_user_idx is None:
        return jsonify({"error": "Inga användarmeddelanden att regenerera"}), 400

    last_user_content = msgs[last_user_idx]["content"]
    thread["messages"]      = msgs[:last_user_idx]
    thread["message_count"] = len(thread["messages"])
    thread["updated_at"]    = _now()
    _persist_save()

    return jsonify({
        "status":            "ready",
        "thread_id":         thread_id,
        "last_user_message": last_user_content,
        "message_count":     thread["message_count"],
        "next_action":       "resend_to_stream",
    })


# ════════════════════════════════════════════════════════════════════════
# [STR-5] /api/threads/search
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/threads/search", methods=["GET"])
def api_threads_search():
    sid = _get_user_id()
    q   = request.args.get("q", "").strip().lower()

    if not q:
        return jsonify({"error": "Sökterm saknas"}), 400
    if len(q) < 2:
        return jsonify({"error": "Söktermen är för kort (minst 2 tecken)"}), 400

    tids    = SESSION_INDEX.get(sid, [])
    results = []

    for tid in tids:
        thread = ALL_THREADS.get(tid)
        if not thread:
            continue

        match_snippets = []
        title_match    = q in thread.get("title", "").lower()

        for m in thread.get("messages", []):
            content_lower = m.get("content", "").lower()
            if q in content_lower:
                idx   = content_lower.find(q)
                start = max(0, idx - 60)
                end   = min(len(m["content"]), idx + len(q) + 60)
                snippet = ("…" if start > 0 else "") + m["content"][start:end] + ("…" if end < len(m["content"]) else "")
                match_snippets.append({"role": m["role"], "snippet": snippet})
                if len(match_snippets) >= 3:
                    break

        if title_match or match_snippets:
            results.append({
                "id":            thread["id"],
                "title":         thread["title"],
                "updated_at":    thread["updated_at"],
                "message_count": thread["message_count"],
                "title_match":   title_match,
                "snippets":      match_snippets,
                "match_count":   len(match_snippets) + (1 if title_match else 0),
            })

    results.sort(key=lambda r: (not r["title_match"], -r["match_count"], r["updated_at"]), reverse=False)
    return jsonify({"query": q, "total": len(results), "results": results})


# ════════════════════════════════════════════════════════════════════════
# [IMG-4] Dedicated image generation endpoint
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/generate-image", methods=["POST"])
def api_generate_image():
    if not DIFFUSERS_OK:
        return jsonify({
            "ok":    False,
            "error": "diffusers inte installerat — kör: pip install diffusers transformers accelerate",
        }), 503

    body            = request.get_json(force=True) or {}
    swedish_prompt  = body.get("prompt",          "").strip()
    negative_prompt = body.get("negative_prompt", "").strip()

    if not swedish_prompt:
        return jsonify({"ok": False, "error": "Fältet 'prompt' saknas"}), 400

    english_prompt = _translate_prompt_to_english(swedish_prompt)
    result         = generate_image_diffusers(english_prompt, negative_prompt)

    if result["ok"]:
        # ── Save to per-user image history ──────────────────────────────────
        sid = _get_user_id()
        import datetime as _dt
        img_record = {
            "id":         uuid.uuid4().hex,
            "b64":        result["b64"],
            "prompt_sv":  swedish_prompt,
            "prompt_en":  english_prompt,
            "created_at": _dt.datetime.utcnow().isoformat() + "Z",
        }
        USER_IMAGES.setdefault(sid, []).insert(0, img_record)
        USER_IMAGES[sid] = USER_IMAGES[sid][:50]   # keep last 50 per user
        _persist_save()                             # ← persist to disk immediately
        # ────────────────────────────────────────────────────────────────────
        return jsonify({
            "ok":        True,
            "b64":       result["b64"],
            "prompt_sv": swedish_prompt,
            "prompt_en": english_prompt,
        })
    else:
        return jsonify({"ok": False, "error": result.get("error", "Okänt fel")}), 503


@app.route("/api/my-images", methods=["GET"])
def api_my_images():
    """Return the current user's generated image history (newest first, max 50)."""
    sid = _get_user_id()
    images = USER_IMAGES.get(sid, [])
    return jsonify({"images": images, "total": len(images)})


# ════════════════════════════════════════════════════════════════════════
# Thread / History management endpoints
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/threads", methods=["GET"])
def api_threads():
    sid        = _get_user_id()
    active_tid = ACTIVE_THREAD.get(sid)
    tids       = SESSION_INDEX.get(sid, [])
    threads_list = sorted(
        [
            {
                "id":            ALL_THREADS[tid]["id"],
                "title":         ALL_THREADS[tid]["title"],
                "created_at":    ALL_THREADS[tid]["created_at"],
                "updated_at":    ALL_THREADS[tid]["updated_at"],
                "message_count": ALL_THREADS[tid]["message_count"],
                "is_active":     tid == active_tid,
                "has_summary":   bool(ALL_THREADS[tid].get("summary")),
            }
            for tid in tids if tid in ALL_THREADS
        ],
        key=lambda x: x["updated_at"],
        reverse=True,
    )
    return jsonify({
        "threads": threads_list, "active_thread": active_tid, "total": len(threads_list),
    })


@app.route("/api/threads/new", methods=["POST"])
def api_new_thread():
    sid    = _get_user_id()
    thread = _make_thread()
    _register_thread(sid, thread)
    return jsonify({
        "thread_id": thread["id"], "thread_title": thread["title"],
        "created_at": thread["created_at"],
    })


@app.route("/api/threads/<thread_id>", methods=["GET"])
def api_get_thread(thread_id: str):
    thread = ALL_THREADS.get(thread_id)
    if not thread:
        return jsonify({"error": "Tråden hittades inte"}), 404
    return jsonify({
        "id": thread["id"], "title": thread["title"],
        "created_at": thread["created_at"], "updated_at": thread["updated_at"],
        "message_count": thread["message_count"], "messages": thread["messages"],
        "summary": thread.get("summary"), "has_summary": bool(thread.get("summary")),
    })


@app.route("/api/threads/<thread_id>", methods=["DELETE"])
def api_delete_thread(thread_id: str):
    sid = _get_user_id()
    if thread_id not in ALL_THREADS:
        return jsonify({"error": "Tråden hittades inte"}), 404

    del ALL_THREADS[thread_id]
    if sid in SESSION_INDEX and thread_id in SESSION_INDEX[sid]:
        SESSION_INDEX[sid].remove(thread_id)

    if ACTIVE_THREAD.get(sid) == thread_id:
        remaining_ids = SESSION_INDEX.get(sid, [])
        remaining = sorted(
            [ALL_THREADS[t] for t in remaining_ids if t in ALL_THREADS],
            key=lambda t: t["updated_at"], reverse=True,
        )
        new_active = remaining[0]["id"] if remaining else None
        ACTIVE_THREAD[sid] = new_active
    else:
        new_active = ACTIVE_THREAD.get(sid)

    _persist_save()
    return jsonify({"status": "deleted", "active_thread": new_active})


@app.route("/api/threads/<thread_id>/switch", methods=["POST"])
def api_switch_thread(thread_id: str):
    sid = _get_user_id()
    if thread_id not in ALL_THREADS:
        return jsonify({"error": "Tråden hittades inte"}), 404
    ACTIVE_THREAD[sid] = thread_id
    thread = ALL_THREADS[thread_id]
    _persist_save()
    return jsonify({
        "thread_id": thread_id, "thread_title": thread["title"],
        "message_count": thread["message_count"],
    })


@app.route("/api/threads/<thread_id>/rename", methods=["POST"])
def api_rename_thread(thread_id: str):
    if thread_id not in ALL_THREADS:
        return jsonify({"error": "Tråden hittades inte"}), 404
    new_title = request.get_json(force=True).get("title", "").strip()
    if not new_title:
        return jsonify({"error": "Titeln får inte vara tom"}), 400
    ALL_THREADS[thread_id]["title"] = new_title[:80]
    _persist_save()
    return jsonify({"status": "ok", "title": new_title[:80]})


# ════════════════════════════════════════════════════════════════════════
# TTS endpoint
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/tts", methods=["POST"])
def api_tts():
    text = request.get_json(force=True).get("text", "")
    path = synth_tts(text)

    if not path or not os.path.exists(path):
        engine = "KBLab Piper" if PIPER_OK else "ingen TTS tillgänglig"
        return jsonify({"error": f"TTS misslyckades ({engine})"}), 503

    if path.endswith(".wav"):
        mp3 = _wav_to_mp3(path)
        if mp3:
            return send_file(mp3, mimetype="audio/mpeg")
        return send_file(path, mimetype="audio/wav")

    return send_file(path, mimetype="audio/mpeg")


@app.route("/api/tts/voices", methods=["GET"])
def api_tts_voices():
    return jsonify({
        "engine":        "kblab-piper" if (PIPER_OK and _piper) else "none",
        "model":         "KBLab/piper-tts-nst-swedish",
        "available":     {"nst-swedish": "KBLab/piper-tts-nst-swedish"},
        "current_voice": "nst-swedish",
    })


# ════════════════════════════════════════════════════════════════════════
# STT endpoint
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/stt/debug", methods=["GET"])
def api_stt_debug():
    import subprocess
    try:
        backends = torchaudio.list_audio_backends() if WAV2VEC_OK else []
    except Exception:
        backends = ["(error)"]

    ffmpeg_ok, ffmpeg_ver = False, ""
    try:
        r = subprocess.run(["ffmpeg", "-version"], capture_output=True, timeout=5)
        ffmpeg_ok  = r.returncode == 0
        ffmpeg_ver = r.stdout.decode().split("\n")[0]
    except Exception as e:
        ffmpeg_ver = str(e)

    pydub_ok = False
    try:
        from pydub import AudioSegment
        pydub_ok = True
    except ImportError:
        pass

    return jsonify({
        "wav2vec_model_loaded": _stt_model is not None,
        "wav2vec_proc_loaded":  _stt_proc  is not None,
        "torchaudio_backends":  backends,
        "ffmpeg_available":     ffmpeg_ok,
        "ffmpeg_version":       ffmpeg_ver,
        "pydub_available":      pydub_ok,
        "recommendation": (
            "OK — all conversion paths available"
            if (ffmpeg_ok and pydub_ok)
            else "Install pydub + ffmpeg: pip install pydub  and add ffmpeg to PATH"
        ),
    })


@app.route("/api/stt", methods=["POST"])
def api_stt():
    if "audio" not in request.files:
        return jsonify({"error": "Ingen ljudfil bifogad"}), 400

    f      = request.files["audio"]
    suffix = os.path.splitext(f.filename or "audio.webm")[1] or ".webm"

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            f.save(tmp.name)
            tmp_path = tmp.name
        size = os.path.getsize(tmp_path)
        if size < 500:
            return jsonify({
                "transcript": "",
                "warning": f"Ljudfilen var för liten ({size} B).",
                "debug": {"file_size_bytes": size},
            })
    except Exception as e:
        return jsonify({"error": f"Kunde inte spara filen: {e}"}), 500

    debug: dict = {"file_size_bytes": size, "suffix": suffix}

    try:
        from pydub import AudioSegment as _AS
        seg = _AS.from_file(tmp_path)
        debug["pydub"] = {
            "ok": True, "duration_ms": len(seg),
            "channels": seg.channels, "frame_rate": seg.frame_rate,
        }
    except Exception as e:
        debug["pydub"] = {"ok": False, "error": str(e)}

    try:
        transcript = kb_transcribe(tmp_path)
        debug["transcript"] = transcript
    except Exception as e:
        _safe_unlink(tmp_path)
        return jsonify({"error": f"Transkribering misslyckades: {e}", "debug": debug}), 500
    finally:
        _safe_unlink(tmp_path)

    if transcript:
        return jsonify({"transcript": transcript, "debug": debug})
    return jsonify({
        "transcript": "",
        "warning": "Ingen text kändes igen — se debug-fältet.",
        "debug": debug,
    })


# ════════════════════════════════════════════════════════════════════════
# Utility endpoints
# ════════════════════════════════════════════════════════════════════════

@app.route("/api/time", methods=["GET"])
def api_time():
    results = tavily_search("vad är klockan i Sverige Stockholm just nu")
    if results:
        direct = next((r["snippet"] for r in results if r["title"] == "Direkt svar"), None)
        answer = direct or results[0].get("snippet", "")
        answer = re.sub(r"#+\s*", "", answer)
        answer = re.sub(r"[*_`]", "", answer).strip()
    else:
        answer = "Kunde inte hämta aktuell tid just nu."
    return jsonify({"reply_sv": answer, "source": "tavily_web_search"})


@app.route("/api/clear", methods=["POST"])
def api_clear():
    sid       = _get_user_id()
    body      = request.get_json(force=True, silent=True) or {}
    thread_id = body.get("thread_id")

    target_tid = thread_id if (thread_id and thread_id in ALL_THREADS) \
                 else ACTIVE_THREAD.get(sid)

    if target_tid and target_tid in ALL_THREADS:
        t = ALL_THREADS[target_tid]
        t["messages"] = []; t["message_count"] = 0
        t["summary"] = None; t["_last_summary_at"] = 0
        t["updated_at"] = _now()
        _persist_save()
        return jsonify({"status": "cleared", "thread_id": target_tid})

    return jsonify({"status": "nothing to clear"})


@app.route("/api/status", methods=["GET"])
def api_status():
    ollama_ok = False
    try:
        ollama_ok = requests.get(f"{OLLAMA_BASE}/api/tags", timeout=3).status_code == 200
    except Exception:
        pass

    sd_loaded = _sd_pipe is not None

    return jsonify({
        "ollama":             ollama_ok,
        "ollama_model":       OLLAMA_MODEL,
        "tts_engine":         "kblab-piper" if (PIPER_OK and _piper) else "none",
        "tts_model":          "KBLab/piper-tts-nst-swedish",
        "tts_ready":          PIPER_OK and _piper is not None,
        "stt_engine":         "kblab-wav2vec2" if (WAV2VEC_OK and _stt_model) else "none",
        "stt_model":          KB_STT_MODEL_ID,
        "stt_ready":          WAV2VEC_OK and _stt_model is not None,
        "kb_stage1_lang":     _tok       is not None,
        "kb_stage2_spell":    _fill_mask is not None,
        "kb_stage3_ner":      _ner_pipe  is not None,
        "kb_stage4_sum":      _bert      is not None,
        "tavily":             bool(TAVILY_KEY),
        "direct_path":        True,
        "direct_categories":  list(DIRECT_ANSWER_TRIGGERS.keys()),
        "thread_management":  True,
        "follow_up_memory":   True,
        "rolling_summary":    True,
        "file_upload":        True,
        "vision_model":       ollama_model_is_vision(),
        "imggen_available":   DIFFUSERS_OK,
        "imggen_model":       SD_MODEL_ID,
        "imggen_pipeline_loaded": sd_loaded,
        "imggen_device":      "cuda" if (TRANSFORMERS_OK and torch.cuda.is_available()) else "cpu",
        "imggen_size":        f"{SD_WIDTH}x{SD_HEIGHT}",
        "imggen_steps":       SD_STEPS,
        "streaming":          True,
        "streaming_endpoint": "/api/chat/stream",
        "stop_generation":    True,
        "edit_message":       True,
        "regenerate":         True,
        "search_chats":       True,
        "fixes_applied": [
            "FIX-1: structured content never truncated",
            "FIX-2: num_predict 512/1024/768",
            "FIX-3: sentence splitter preserves lists",
            "FIX-4: direct_path max_sents=5",
            "FIX-5: tech terms skipped in spell-check",
            "FIX-6: markdown preserved in chat reply",
            "UP-1:  /api/chat accepts multipart/form-data + JSON",
            "UP-2:  file text extracted → injected into Ollama prompt",
            "UP-3:  images base64-encoded for vision models",
            "UP-4:  temp files cleaned up after every response",
            "IMG-FIX-1: is_image_generation_request uses Swedish creation verbs only",
            "IMG-FIX-2: image generation check runs BEFORE Tavily search",
            "IMG-FIX-3: _translate_prompt_to_english enabled and calls Ollama",
            "IMG-FIX-4: English input rejected with Swedish-only message",
            "STR-1: /api/chat/stream SSE streaming endpoint",
            "STR-2: /api/chat/stop  cancel in-progress generation",
            "STR-3: /api/threads/<id>/messages/<idx>/edit  edit & resend",
            "STR-4: /api/threads/<id>/regenerate  regenerate last reply",
            "STR-5: /api/threads/search  full-text search across chats",
        ],
    })


@app.route("/api/models", methods=["GET"])
def api_models():
    try:
        r = requests.get(f"{OLLAMA_BASE}/api/tags", timeout=5)
        r.raise_for_status()
        models = r.json().get("models", [])
        result = []
        for m in models:
            name       = m.get("name", "")
            size_bytes = m.get("size", 0)
            size_gb    = round(size_bytes / 1e9, 1) if size_bytes else None
            is_vision  = any(v in name.lower() for v in _VISION_MODEL_NAMES)
            result.append({
                "name": name, "size_gb": size_gb, "is_vision": is_vision,
                "is_active": name == OLLAMA_MODEL or name.split(":")[0] == OLLAMA_MODEL.split(":")[0],
            })
        result.sort(key=lambda x: (not x["is_active"], x["name"]))
        return jsonify({"models": result, "active": OLLAMA_MODEL})
    except Exception as e:
        return jsonify({"models": [], "active": OLLAMA_MODEL, "error": str(e)})


@app.route("/api/models/select", methods=["POST"])
def api_models_select():
    global OLLAMA_MODEL
    name = (request.get_json(force=True) or {}).get("model", "").strip()
    if not name:
        return jsonify({"error": "Modellnamn saknas"}), 400
    OLLAMA_MODEL = name
    logger.info(f"[Models] Switched active model → {OLLAMA_MODEL}")
    return jsonify({"active": OLLAMA_MODEL, "is_vision": ollama_model_is_vision()})


@app.route("/api/tavily-test", methods=["GET"])
def api_tavily_test():
    q = request.args.get("q", "vad är vädret i Stockholm idag")
    results = tavily_search(q)
    return jsonify({
        "query": q,
        "key_present": bool(TAVILY_KEY),
        "results_count": len(results),
        "results": results
    })


if __name__ == "__main__":
    print("""
╔══════════════════════════════════════════════════════════════════════╗
║  🇸🇪  Swedish AI Chat  —  KB-Lab Full Stack  (v6 — ImgGen Fixed)   ║
╠══════════════════════════════════════════════════════════════════════╣
║  IMAGE GENERATION FIXES (v6)                                         ║
║   IMG-FIX-1  is_image_generation_request: Swedish verbs only        ║
║   IMG-FIX-2  Image check runs BEFORE Tavily (no interception)       ║
║   IMG-FIX-3  Translation to English via Ollama — ENABLED            ║
║   IMG-FIX-4  English input → rejected with Swedish-only message     ║
╠══════════════════════════════════════════════════════════════════════╣
║  Example Swedish triggers:                                           ║
║   "rita en röd bil i solnedgången"   → generates image              ║
║   "skapa en bild av ett berg"        → generates image              ║
║   "generera en bild av en katt"      → generates image              ║
║   "kan du måla ett landskap"         → generates image              ║
║   "en röd bil i solnedgången"        → normal chat (no verb)        ║
╚══════════════════════════════════════════════════════════════════════╝
""")
    app.run(debug=True, host="0.0.0.0", port=5000, threaded=True)